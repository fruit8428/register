#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
扶輪 3523 地區 2026-27 年度保齡球聯誼賽
郵件自動化對帳與報名統計系統 (Bowling Assistant)

功能：
1. 自動連線 ssrotary@ms41.hinet.net 信箱。
2. 第一次執行/點擊時：自動往回推算 3 天郵件進行檢測。
3. 此後每 3 小時自動背景偵測；人工點擊時自動回溯到上一次偵測時間點。
4. 精準辨識「上海商業儲蓄銀行」入帳通知 Mail，抓取匯款日期、金額、社名並填入「報名統計.xlsx」。
5. 精準辨識各社執秘寄送之「保齡球聯誼賽 (附件二)」報名表（Word/Excel/郵件內文），解析參賽人員名單、用餐名單與主協辦狀態。
6. 將報名人數與費用統計自動同步至「報名統計.xlsx」。
7. 將各社參賽名單依社名為單位依序排列，自動累積更新至「附件四.xlsx」（參賽名單）。
8. 自動備份資料、記錄已處理信件，支援單次執行、定時背景監聽與即時報表輸出。
"""

import os
import re
import sys
import json
import ssl
import time
import shutil
import hashlib
import imaplib
import email
import email.utils
import tempfile
import subprocess
import io
import unicodedata
from email.header import decode_header
from datetime import datetime, timedelta, timezone
import openpyxl

# ==================== 設定檔 (Configuration) ====================
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
CLUBS_DB_PATH = os.path.join(SCRIPT_DIR, "clubs_db.json")
STATS_XLSX_PATH = os.path.join(SCRIPT_DIR, "報名統計.xlsx")
ROSTER_XLSX_PATH = os.path.join(SCRIPT_DIR, "附件四.xlsx")

DATA_DIR = os.path.join(SCRIPT_DIR, "data")
BACKUP_DIR = os.path.join(SCRIPT_DIR, "backups")
PROCESSED_FILE = os.path.join(DATA_DIR, "processed_emails.json")
ROSTER_DATA_FILE = os.path.join(DATA_DIR, "roster_data.json")
PENDING_FILE = os.path.join(DATA_DIR, "pending_payments.json")
SYNC_STATE_FILE = os.path.join(DATA_DIR, "sync_state.json")

# 信箱設定
IMAP_SERVER = "ms41.hinet.net"
IMAP_PORT = 993
MAIL_USER = "ssrotary"
MAIL_PASS = "1688@Ss88"

# 銀行帳號設定
BANK_NAME = "上海商業儲蓄銀行"
BANK_BRANCH = "南京東路分行"
TARGET_ACCOUNT = "40102000039270"
TARGET_ACCOUNT_LAST5 = "39270"

# 自動排程週期（預設 3 小時）
AUTO_SYNC_INTERVAL_HOURS = 3

# 台北本地時區 (UTC+8)
LOCAL_TZ = timezone(timedelta(hours=8))

os.makedirs(DATA_DIR, exist_ok=True)
os.makedirs(BACKUP_DIR, exist_ok=True)


# ==================== 同步狀態管理 (Sync State) ====================

def get_sync_state():
    """取得上次同步與排程狀態"""
    if os.path.exists(SYNC_STATE_FILE):
        try:
            with open(SYNC_STATE_FILE, "r", encoding="utf-8") as f:
                return json.load(f)
        except:
            pass
    return {
        "last_sync_time": None,
        "next_scheduled_sync": None,
        "interval_hours": AUTO_SYNC_INTERVAL_HOURS,
        "is_first_sync": True,
        "total_sync_count": 0
    }


def save_sync_state(state):
    """儲存同步狀態"""
    with open(SYNC_STATE_FILE, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False, indent=2)


# ==================== 輔助工具函式 ====================

def decode_mime_words(s):
    """解碼 MIME 編碼的信件標題或文字"""
    if not s:
        return ""
    dh = decode_header(s)
    res = []
    for text, enc in dh:
        if isinstance(text, bytes):
            res.append(text.decode(enc or "utf-8", errors="ignore"))
        else:
            res.append(str(text))
    return "".join(res).strip()


def backup_file(filepath):
    """備份檔案至 backups/ 目錄"""
    if not os.path.exists(filepath):
        return
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base_name = os.path.basename(filepath)
    name, ext = os.path.splitext(base_name)
    backup_path = os.path.join(BACKUP_DIR, f"{name}_{timestamp}{ext}")
    shutil.copy2(filepath, backup_path)
    return backup_path


def load_clubs_db():
    """載入 80 社基本資料庫"""
    if not os.path.exists(CLUBS_DB_PATH):
        raise FileNotFoundError(f"找不到社名資料庫: {CLUBS_DB_PATH}")
    with open(CLUBS_DB_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


def match_club_from_text(text, clubs):
    """
    從文字中比對扶輪社：
    1. 完整社名 (例如: 台北市西南區扶輪社)
    2. 標籤 (例如: 12-1龍門社)
    3. 編號 (例如: 12-1)
    4. 簡稱 (例如: 龍門社)
    5. 核心特色根詞 (例如: 龍門, 逸仙, 西南區)
    """
    if not text:
        return None

    # 過濾常見系統與公用干擾詞，避免誤匹配
    filtered_text = text
    filtered_text = re.sub(r"網路銀行|電子郵件服務|網址|網路連線", "網銀", filtered_text)
    filtered_text = re.sub(r"台北市松山區南京東路[^\n\r]*|松山區南京東路[^\n\r]*", "", filtered_text)
    filtered_text = re.sub(r"台北松山扶輪社林永曄|40102000039270|40102\*{7}70|ssrotary@ms41\.hinet\.net", "", filtered_text)

    cleaned_text = re.sub(r"[\s\-_　]", "", filtered_text)

    # 1. 完整社名比對
    for c in sorted(clubs, key=lambda x: len(x["name"]), reverse=True):
        if c["name"] in filtered_text or c["name"].replace("臺", "台") in filtered_text:
            return c

    # 2. 標籤比對 (如 12-1龍門社)
    for c in clubs:
        if c["tag"] in filtered_text or c["tag"] in cleaned_text:
            return c

    # 3. 編號比對 (如 12-1, 01-1)
    code_match = re.search(r"\b(0?[1-9]|1[0-3])-(0?[1-9])\b", filtered_text)
    if code_match:
        norm_code = f"{int(code_match.group(1)):02d}-{int(code_match.group(2))}"
        for c in clubs:
            if c["code"] == norm_code:
                return c

    # 4. 簡稱比對 (如 龍門社)
    for c in sorted(clubs, key=lambda x: len(x["short"]), reverse=True):
        if c["short"] in filtered_text:
            return c

    # 5. 核心詞比對 (如 龍門, 逸仙)
    for c in sorted(clubs, key=lambda x: len(x["short"][:-1]), reverse=True):
        root = c["short"][:-1]
        # 排除過短或容易與日常詞彙衝突之根詞
        if root in ["網路"]:
            continue
        if len(root) >= 2 and root in filtered_text:
            return c

    return None


# ==================== 郵件內文與附件提取 ====================

def extract_email_body_and_attachments(msg):
    """解析郵件純文字與附件"""
    body_text = ""
    attachments = []

    for part in msg.walk():
        content_type = part.get_content_type()
        content_disposition = str(part.get("Content-Disposition", ""))
        filename = part.get_filename()

        if filename:
            decoded_filename = decode_mime_words(filename)
            payload = part.get_payload(decode=True)
            if payload:
                attachments.append({
                    "filename": decoded_filename,
                    "data": payload
                })
        elif content_type == "text/plain" and "attachment" not in content_disposition:
            charset = part.get_content_charset() or "utf-8"
            payload = part.get_payload(decode=True)
            if payload:
                body_text += "\n" + payload.decode(charset, errors="ignore")
        elif content_type == "text/html" and "attachment" not in content_disposition:
            charset = part.get_content_charset() or "utf-8"
            payload = part.get_payload(decode=True)
            if payload and not body_text.strip():
                from bs4 import BeautifulSoup
                soup = BeautifulSoup(payload.decode(charset, errors="ignore"), "html.parser")
                body_text += "\n" + soup.get_text(separator="\n")

    return body_text.strip(), attachments


# ==================== 匯款水單圖片 OCR 自動辨識 ====================

def ocr_image_slip(img_bytes):
    """
    辨識匯款/轉帳水單截圖：
    利用 macOS 原生 Vision 框架 (或 tesseract 備援) 辨識圖片中的文字，
    精準提取：交易金額、交易日期、付款帳號末五碼、備註。
    """
    if not img_bytes:
        return {}

    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tf:
        tf.write(img_bytes)
        tpath = tf.name

    ocr_text = ""
    # 優先使用 macOS 原生 Vision 框架 (繁體中文高精度辨識)
    swift_code = f"""
    import Cocoa
    import Vision

    let imageURL = URL(fileURLWithPath: "{tpath}")
    guard let image = NSImage(contentsOf: imageURL),
          let cgImage = image.cgImage(forProposedRect: nil, context: nil, hints: nil) else {{
        exit(1)
    }}

    let request = VNRecognizeTextRequest {{ (request, error) in
        guard let observations = request.results as? [VNRecognizedTextObservation] else {{ return }}
        for observation in observations {{
            let topCandidate = observation.topCandidates(1)
            if let recognizedText = topCandidate.first {{
                print(recognizedText.string)
            }}
        }}
    }}
    request.recognitionLanguages = ["zh-Hant", "en-US"]
    request.recognitionLevel = .accurate

    let handler = VNImageRequestHandler(cgImage: cgImage, options: [:])
    try? handler.perform([request])
    """
    try:
        res = subprocess.run(["swift", "-e", swift_code], capture_output=True, text=True, timeout=10)
        if res.returncode == 0 and res.stdout.strip():
            ocr_text = res.stdout.strip()
    except Exception:
        pass

    # 備援：若無 Vision 則嘗試 tesseract
    if not ocr_text and shutil.which("tesseract"):
        try:
            t_res = subprocess.run(["tesseract", tpath, "stdout"], capture_output=True, text=True, timeout=10)
            if t_res.returncode == 0:
                ocr_text = t_res.stdout.strip()
        except Exception:
            pass

    try:
        os.remove(tpath)
    except Exception:
        pass

    if not ocr_text:
        return {}

    slip_info = {
        "amount": None,
        "date": None,
        "last5": None,
        "memo": "",
        "raw_text": ocr_text
    }

    # 1. 金額 (例如: TWD 13,400 或 NT$ 13,400 或 金額: 13400)
    amt_m = re.search(r"(?:TWD|NT\$?|\$|金額)[\s:]*([0-9,]+)", ocr_text, re.I)
    if not amt_m:
        amt_m = re.search(r"([0-9,]+)\s*元", ocr_text)
    if amt_m:
        try:
            slip_info["amount"] = int(amt_m.group(1).replace(",", "").strip())
        except ValueError:
            pass

    # 2. 日期 (例如: 2026/09/24 11:41:13)
    date_m = re.search(r"(\d{4}[-/年]\d{1,2}[-/月]\d{1,2})", ocr_text)
    if date_m:
        dstr = date_m.group(1).replace("年", "/").replace("月", "/").replace("-", "/")
        parts = [int(p) for p in dstr.split("/") if p.strip()]
        if len(parts) == 3:
            slip_info["date"] = f"{parts[0]:04d}/{parts[1]:02d}/{parts[2]:02d}"

    # 3. 轉出帳號末五碼
    # 付款帳號 008 - 1062006533710016
    for line in ocr_text.splitlines():
        clean_line = re.sub(r"[\s\-]", "", line)
        m_acc = re.search(r"(\d{10,20})", clean_line)
        if m_acc:
            acc = m_acc.group(1)
            # 排除目標主辦專戶 (40102000039270)
            if not acc.startswith("40102000039270") and not acc.endswith("39270"):
                slip_info["last5"] = acc[-5:]
                break

    # 4. 備註 / 留言
    memo_m = re.search(r"(?:留言給對方|給自己備忘錄|備註|附言)[：:\s]*([^\n\r]+)", ocr_text)
    if memo_m:
        slip_info["memo"] = memo_m.group(1).strip()

    return slip_info


# ==================== 上海商銀入帳通知精準解析 ====================

def parse_scsb_deposit_email(subject, body, clubs):
    """
    精準解析上海商業儲蓄銀行「帳戶入帳通知(E-MAIL通知)」
    排除對帳單、行銷廣告、請款授權等無關信件。
    """
    exclude_kws = ["對帳單", "批次授權", "申請書", "優利定存", "佳節快樂", "驗證信", "重要訊息提醒"]
    if any(k in subject for k in exclude_kws):
        return None

    is_deposit_subject = any(k in subject for k in ["入帳通知", "存款入帳", "轉帳入帳", "匯入款項", "帳戶入帳"])
    is_deposit_body = any(k in body for k in ["入帳通知", "存款入帳通知", "臺幣存款帳戶入帳通知"])

    if not (is_deposit_subject or is_deposit_body):
        return None

    result = {
        "type": "scsb_deposit",
        "date": None,
        "amount": None,
        "remitter": "",
        "memo": "",
        "account_last5": "",
        "club": None,
        "is_encrypted_notice": False,
        "raw_subject": subject
    }

    # 1. 日期 (先檢查民國年，如 115年09月24日)
    roc_match = re.search(r"(11\d)[-/年](\d{1,2})[-/月](\d{1,2})", subject + " " + body)
    if roc_match:
        y = int(roc_match.group(1)) + 1911
        m = int(roc_match.group(2))
        d = int(roc_match.group(3))
        result["date"] = f"{y:04d}/{m:02d}/{d:02d}"
    else:
        date_match = re.search(r"(\d{4}[-/年]\d{1,2}[-/月]\d{1,2})", body)
        if date_match:
            dstr = date_match.group(1).replace("年", "/").replace("月", "/").replace("-", "/")
            parts = [int(p) for p in dstr.split("/") if p.strip()]
            if len(parts) == 3:
                result["date"] = f"{parts[0]:04d}/{parts[1]:02d}/{parts[2]:02d}"

    if not result["date"]:
        result["date"] = datetime.now().strftime("%Y/%m/%d")

    # 2. 金額
    amt_match = re.search(r"(?:入帳金額|交易金額|金額)[：:\s]*(?:新臺幣|NT\$?)?\s*([0-9,]+)", body)
    if not amt_match:
        amt_match = re.search(r"(?:新臺幣|NT\$)\s*([0-9,]+)\s*元", body)

    if amt_match:
        amt_str = amt_match.group(1).replace(",", "").strip()
        try:
            result["amount"] = int(amt_str)
        except ValueError:
            pass
    elif "加密型態提供" in body or "解密開啟附件" in body:
        result["is_encrypted_notice"] = True

    # 3. 備註 / 存入人 / 匯款人 / 末五碼
    memo_match = re.search(r"(?:備註|摘要|存入人|匯款人|附言|轉出帳號)[：:\s]*([^\n\r]+)", body)
    if memo_match:
        result["memo"] = memo_match.group(1).strip()

    last5_match = re.search(r"(?:末[四五]碼|帳號末[四五]碼|末\s*5\s*碼)[：:\s]*([0-9]{4,5})", body)
    if last5_match:
        result["account_last5"] = last5_match.group(1)

    # 4. 比對扶輪社 (過濾銀行制式文字干擾)
    clean_body_for_match = re.sub(r"網路銀行|重要訊息提醒|電子郵件服務|松山區|南京東路|對帳單", "", body)
    matched_club = match_club_from_text(subject + " " + clean_body_for_match, clubs)
    result["club"] = matched_club

    return result


# ==================== 附件二保齡球報名表解析 ====================

def is_bowling_registration_email(subject, attachments, body):
    """判斷信件是否與本地區保齡球聯誼賽報名相關"""
    exclude_events = ["登山", "健行", "高爾夫", "GOLF", "音樂會", "研討會", "小兒麻痺", "泰國", "水上", "交換甄選", "敬老", "就職", "對帳單"]
    if any(k in subject for k in exclude_events) and "保齡" not in subject:
        return False

    has_bowling = ("保齡" in subject) or any("保齡" in a["filename"] for a in attachments)
    if not has_bowling and any("附件二" in a["filename"] for a in attachments):
        has_bowling = ("保齡" in body) or ("保齡" in subject)

    return has_bowling


def parse_docx_bytes(docx_bytes):
    """解析 Word (.docx) 附件二報名表"""
    import docx
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as tmp:
        tmp.write(docx_bytes)
        tmp_path = tmp.name

    try:
        doc = docx.Document(tmp_path)
        data = {
            "zone": "",
            "club_name": "",
            "contact": "",
            "phone": "",
            "sponsor": None,
            "players": [],
            "lunch_only": []
        }

        full_text = "\n".join([p.text.strip() for p in doc.paragraphs if p.text.strip()])
        m_info = re.search(r"(\d+)?\s*分區\s*([^\s　]+)?\s*扶輪社.*?聯絡人[：:\s]*([^\s　]+)?.*?行動電話[：:\s]*([^\s　]+)?", full_text)
        if m_info:
            data["zone"] = m_info.group(1) or ""
            data["club_name"] = m_info.group(2) or ""
            data["contact"] = m_info.group(3) or ""
            data["phone"] = m_info.group(4) or ""

        for table in doc.tables:
            headers = [c.text.strip().replace("\n", " ") for c in table.rows[0].cells]
            if any("姓名" in h or "姓  名" in h for h in headers) and any("參加" in h for h in headers):
                for row in table.rows[1:]:
                    vals = [c.text.strip().replace("\n", " ") for c in row.cells]
                    if not vals or vals[0] in ["例", "範例"]:
                        continue
                    name = vals[1] if len(vals) > 1 else ""
                    nick = vals[2] if len(vals) > 2 else ""
                    gender = vals[3] if len(vals) > 3 else ""
                    phone = vals[4] if len(vals) > 4 else ""
                    spouse = vals[5] if len(vals) > 5 else ""
                    child = vals[6] if len(vals) > 6 else ""
                    veg = vals[7] if len(vals) > 7 else ""
                    category = vals[8] if len(vals) > 8 else ""

                    if not name and not nick:
                        continue

                    gender_val = "女" if "女" in gender else "男"
                    cat_upper = category.upper()

                    if "A" in cat_upper or "參" in category or (not cat_upper and name):
                        data["players"].append({
                            "name": name,
                            "nickname": nick,
                            "gender": gender_val,
                            "phone": phone,
                            "is_spouse": bool(spouse and spouse != ""),
                            "is_child": bool(child and child != ""),
                            "is_veg": bool(veg and veg != "")
                        })
                    elif "B" in cat_upper or "餐" in category:
                        data["lunch_only"].append({
                            "name": name,
                            "nickname": nick,
                            "gender": gender_val,
                            "phone": phone,
                            "is_veg": bool(veg and veg != "")
                        })

            for row in table.rows:
                rtxt = " ".join([c.text for c in row.cells])
                if "共同主辦社" in rtxt and any(mark in rtxt for mark in ["■", "√", "V", "v", "★"]):
                    data["sponsor"] = "主辦社"
                elif "協辦社" in rtxt and any(mark in rtxt for mark in ["■", "√", "V", "v", "★"]):
                    data["sponsor"] = "協辦社"
                elif "贊助社" in rtxt and any(mark in rtxt for mark in ["■", "√", "V", "v", "★"]):
                    data["sponsor"] = "贊助社"

        return data
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)


def parse_doc_bytes(doc_bytes):
    """解析舊版 Word (.doc) 格式"""
    with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as tmp:
        tmp.write(doc_bytes)
        tmp_path = tmp.name

    txt_path = tmp_path.replace(".doc", ".txt")
    try:
        subprocess.run(["/usr/bin/textutil", "-convert", "txt", tmp_path, "-output", txt_path], check=True)
        with open(txt_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
        return parse_text_registration(text)
    except Exception as e:
        print(f"    [解析 .doc 提示]: {e}")
        return None
    finally:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)
        if os.path.exists(txt_path):
            os.remove(txt_path)


def parse_pdf_bytes(pdf_bytes):
    """解析 PDF 格式的報名表附件"""
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(pdf_bytes))
        full_text = ""
        for page in reader.pages:
            full_text += (page.extract_text() or "") + "\n"
        return parse_text_registration(full_text)
    except Exception as e:
        print(f"    [解析 .pdf 提示]: {e}")
        return None


def is_header_or_noise(text):
    """判斷該文字是否為報名表表頭、欄位標題或說明性雜訊"""
    t = text.strip()
    clean_t = re.sub(r"[\s\-_　\:\：\︰\.]", "", t)
    noise_phrases = [
        "附件", "國際扶輪", "保齡球", "聯誼賽", "報名表", "分區", "扶輪社", "聯絡人", "行動電話",
        "報名說明", "匯款", "銀行", "戶名", "帳號", "序號", "nickname", "性別", "姓名",
        "電話", "寶尊眷", "寶尊", "社員", "子女", "社員子女", "素食", "小計", "合計", "共同主辦", "協辦社", "贊助社",
        "以利完成", "e-mail", "tel", "範例", "參加a", "參賽b", "僅午餐", "報名費", "餐費", "費用", "1,200", "1200", "800", "元×", "元x", "繳費"
    ]
    if any(nw in clean_t.lower() for nw in noise_phrases):
        return True
    if clean_t.lower() in ["寶", "尊", "眷", "社", "員", "子", "女", "素", "食", "男", "女", "參加", "參賽", "僅午餐", "a", "b", "性", "別"]:
        return True
    return False


def parse_single_player_line(line):
    """
    解析單行參賽者名單，支援複合職稱/英文名/暱稱
    例如: PP Sure 男, P Lin.C夫人 女, PP Borker 男, Bank 男, 謝淮哲 PE Eric 男 0938-158058 A
    """
    line = line.strip()
    if not line or is_header_or_noise(line):
        return None

    # 去除序號，例如 "1.", "1 ", "1、", "#1", "(1)"
    line = re.sub(r"^\s*(?:#?\d+[\.\、\s\-\:\)]+|\(\d+\))\s*", "", line).strip()
    if not line or line.isdigit() or is_header_or_noise(line):
        return None

    # 檢查末尾是否有分類 A（參賽）或 B（僅用餐）
    category = "A"
    cat_match = re.search(r"[,\s\t]+([ABab]|[參用][賽餐]?)\s*$", line)
    if cat_match:
        cat_token = cat_match.group(1).upper()
        if cat_token in ["B", "用", "用餐", "僅用餐"]:
            category = "B"
        line = line[:cat_match.start()].strip()

    # 匹配性別 (男 / 女 / M / F / Male / Female)
    m = re.search(r"^(.*?)[,\s\t]+([男女]|[Mm]ale|[Ff]emale|[MFmf])(?:\s+[\d\-]+)?(?:\s+.*)?$", line)
    if m:
        full_name_part = m.group(1).strip()
        g_raw = m.group(2).strip().lower()
        gender = "女" if g_raw in ["女", "f", "female"] else "男"
    else:
        # 如果整行完全沒有性別標記
        if len(line) > 15 or len(line) < 2 or any(punc in line for punc in ["：", ":", "。", "、", "(", ")", "（", "）", "！", "!"]):
            return None
        full_name_part = line
        gender = "女" if "夫人" in line else "男"

    full_name_part = re.sub(r"^[\,\，\s]+|[\,\，\s]+$", "", full_name_part).strip()
    if not full_name_part or full_name_part.isdigit() or is_header_or_noise(full_name_part):
        return None

    # 若有 中文姓名 + 暱稱 (例如: "謝淮哲 PE Eric")
    # 但若開頭是扶輪職稱 (例如 "PP Sure", "P Lin.C夫人", "PP Borker")，則為整體姓名，不予拆分
    double_name_match = re.match(r"^([\u4e00-\u9fa5]{2,4})\s+([A-Za-z0-9\.\s\u4e00-\u9fa5]+)$", full_name_part)
    if double_name_match and not any(full_name_part.startswith(p) for p in ["PP", "P ", "CP", "IPP", "PE", "PN", "VP", "AG", "DG", "PDG"]):
        real_name = double_name_match.group(1).strip()
        nick_name = double_name_match.group(2).strip()
    else:
        real_name = full_name_part
        nick_name = full_name_part

    return {
        "name": real_name,
        "nickname": nick_name,
        "gender": gender,
        "category": category
    }


def parse_text_registration(text):
    """從純文字內容（含 PDF 提取文字）提取報名資料"""
    if not text:
        return None

    text = unicodedata.normalize("NFKC", text)
    text = re.sub(r"([男女])(\d)", r"\1 \2", text)

    data = {
        "zone": "",
        "club_name": "",
        "contact": "",
        "phone": "",
        "sponsor": None,
        "total_amount": None,
        "players": [],
        "lunch_only": []
    }

    # 正規化聯絡人與電話標頭
    clean_header_text = re.sub(r"(聯絡人|行動電話|電話)[\.\:\s︰]+", r"\1: ", text)
    m_info = re.search(r"([0-9一二三四五六七八九十]+)?\s*分區\s*([^\s　]+)?\s*扶輪社.*?聯絡人:\s*(.*?)\s+行動電話:\s*([0-9\-]+)", clean_header_text)
    if m_info:
        data["zone"] = m_info.group(1) or ""
        data["club_name"] = m_info.group(2) or ""
        data["contact"] = m_info.group(3).strip() or ""
        data["phone"] = m_info.group(4).strip() or ""

    # 贊助判斷
    if "共同主辦社" in text and any(m in text for m in ["■ 共同主辦", "√ 共同主辦", "V 共同主辦", "■共同主辦", "★共同主辦", "★ 共同主辦"]):
        data["sponsor"] = "主辦社"
    elif "協辦社" in text and any(m in text for m in ["■ 協辦", "√ 協辦", "V 協辦", "■協辦", "★協辦", "★ 協辦"]):
        data["sponsor"] = "協辦社"
    elif "贊助社" in text and any(m in text for m in ["■ 贊助", "√ 贊助", "V 贊助", "■贊助", "★贊助", "★ 贊助"]):
        data["sponsor"] = "贊助社"

    # 若未直接勾選，依總額差異計算贊助類別 (差額 10000 為主辦社，5000 為協辦社)
    m_total = re.search(r"合計[：:\s]*([\d,]+)", text)
    if m_total:
        try:
            data["total_amount"] = int(m_total.group(1).replace(",", "").strip())
        except ValueError:
            pass

    if not data["sponsor"] and data["total_amount"]:
        m_a = re.search(r"參賽人員報名費.*?小計[：:\s]*([\d,]+)", text)
        m_b = re.search(r"只參加午餐會.*?小計[：:\s]*([\d,]+)", text)
        fee_a = int(m_a.group(1).replace(",", "")) if m_a and m_a.group(1) else 0
        fee_b = int(m_b.group(1).replace(",", "")) if m_b and m_b.group(1) else 0
        diff = data["total_amount"] - (fee_a + fee_b)
        if diff == 10000:
            data["sponsor"] = "主辦社"
        elif diff == 5000:
            data["sponsor"] = "協辦社"
        elif diff == 3000:
            data["sponsor"] = "贊助社"

    lines = text.splitlines()
    for line in lines:
        p = parse_single_player_line(line)
        if p:
            if p["category"] == "B":
                data["lunch_only"].append({
                    "name": p["name"],
                    "nickname": p["nickname"],
                    "gender": p["gender"]
                })
            else:
                data["players"].append({
                    "name": p["name"],
                    "nickname": p["nickname"],
                    "gender": p["gender"]
                })

    return data


# ==================== Excel 更新操作 ====================

def update_registration_stats_excel(club, sponsor=None, players_count=None, lunch_count=None,
                                   remit_date=None, remit_amt=None, note=None):
    """
    更新「報名統計.xlsx」對應扶輪社列，完整保留既有公式
    Col 4 (D): 共同主辦社 / 協辦社
    Col 7 (G): 參賽人數
    Col 9 (I): 僅用餐不參賽人數
    Col 12 (L): 匯款日期
    Col 13 (M): 收款金額
    Col 15 (O): 備註
    """
    backup_file(STATS_XLSX_PATH)
    wb = openpyxl.load_workbook(STATS_XLSX_PATH, data_only=False)
    sheet = wb["各社報名資料"]
    row_idx = club["row"]

    if sponsor is not None:
        sheet.cell(row_idx, 4).value = sponsor
    if players_count is not None:
        sheet.cell(row_idx, 7).value = int(players_count)
    if lunch_count is not None:
        sheet.cell(row_idx, 9).value = int(lunch_count)
    if remit_date is not None:
        sheet.cell(row_idx, 12).value = str(remit_date)
    if remit_amt is not None:
        sheet.cell(row_idx, 13).value = int(remit_amt)
    if note is not None:
        cur_note = sheet.cell(row_idx, 15).value or ""
        if cur_note and note not in str(cur_note):
            sheet.cell(row_idx, 15).value = f"{cur_note}；{note}"
        else:
            sheet.cell(row_idx, 15).value = note

    wb.save(STATS_XLSX_PATH)
    print(f"  [Excel更新] 報名統計.xlsx 列 {row_idx} ({club['code']} {club['name']}) 已更新！")


def update_roster_excel(all_registered_data, clubs):
    """
    更新「附件四.xlsx」（參賽名單）：
    將所有已報名社的 Category A 選手填入 40 個球道（每道 4 名，共 160 名額）
    若選手先前已被人工調整賽道，保留其已指定之賽道位置；
    新報名選手則依序填入尚未被佔用的空席位。
    """
    backup_file(ROSTER_XLSX_PATH)
    wb = openpyxl.load_workbook(ROSTER_XLSX_PATH)
    sheet = wb["參賽名單"]

    # 1. 整理所有已報名選手名單
    all_players = []
    sorted_clubs = sorted(clubs, key=lambda c: c["code"])

    for c in sorted_clubs:
        c_code = c["code"]
        if c_code in all_registered_data:
            c_info = all_registered_data[c_code]
            players = c_info.get("players", [])
            for p in players:
                display_name = p.get("nickname") or p.get("name")
                p_gender = p.get("gender", "男")
                if p_gender not in ["男", "女"]:
                    if any(x in str(p_gender).lower() for x in ["女", "f", "female"]):
                        p_gender = "女"
                    elif "夫人" in str(display_name):
                        p_gender = "女"
                    else:
                        p_gender = "男"
                all_players.append({
                    "name": display_name,
                    "gender": p_gender,
                    "club_tag": c["tag"]
                })

    print(f"  [附件四彙整] 累計參賽選手總數: {len(all_players)} / 160 人")

    # 2. 讀取現有 160 席位目前排定情況
    current_roster = [None] * 160
    for slot_idx in range(160):
        row_idx = slot_idx + 2
        p_name = sheet.cell(row_idx, 3).value
        p_gender = sheet.cell(row_idx, 4).value
        p_club = sheet.cell(row_idx, 5).value
        if p_name:
            current_roster[slot_idx] = {
                "name": str(p_name),
                "gender": str(p_gender) if p_gender else "男",
                "club_tag": str(p_club) if p_club else ""
            }

    # 3. 找出仍有效的既有選手並保留位置，移除已不在報名資料中的選手
    assigned_slots = [None] * 160
    remaining_players = list(all_players)

    for slot_idx, slot in enumerate(current_roster):
        if slot:
            match_idx = -1
            for p_idx, p in enumerate(remaining_players):
                if p["name"] == slot["name"] and p["club_tag"] == slot["club_tag"]:
                    match_idx = p_idx
                    break
            if match_idx != -1:
                assigned_slots[slot_idx] = remaining_players.pop(match_idx)

    # 4. 將尚未排入的全新選手，依序填入剩餘空位
    for p in remaining_players:
        for slot_idx in range(160):
            if assigned_slots[slot_idx] is None:
                assigned_slots[slot_idx] = p
                break

    # 5. 寫回 Excel (Row 2 到 Row 161)
    for slot_idx in range(160):
        row_idx = slot_idx + 2
        lane = (slot_idx // 4) + 1
        seq = (slot_idx % 4) + 1

        sheet.cell(row_idx, 1).value = lane
        sheet.cell(row_idx, 2).value = seq

        p = assigned_slots[slot_idx]
        if p:
            sheet.cell(row_idx, 3).value = p["name"]
            sheet.cell(row_idx, 4).value = p["gender"]
            sheet.cell(row_idx, 5).value = p["club_tag"]
        else:
            sheet.cell(row_idx, 3).value = None
            sheet.cell(row_idx, 4).value = None
            sheet.cell(row_idx, 5).value = None

    wb.save(ROSTER_XLSX_PATH)
    print(f"  [Excel更新] 附件四.xlsx 參賽名單已更新完畢（已保留自訂賽道排位）！")


def swap_players_in_roster(from_lane, from_seq, to_lane, to_seq):
    """
    在「附件四.xlsx」中移動或互換選手賽道席位
    from_lane, to_lane: 1~40
    from_seq, to_seq: 1~4
    """
    from_lane = int(from_lane)
    from_seq = int(from_seq)
    to_lane = int(to_lane)
    to_seq = int(to_seq)

    if not (1 <= from_lane <= 40 and 1 <= from_seq <= 4 and 1 <= to_lane <= 40 and 1 <= to_seq <= 4):
        raise ValueError("球道號碼需介於 1~40，席位序號需介於 1~4")

    from_row = (from_lane - 1) * 4 + from_seq + 1
    to_row = (to_lane - 1) * 4 + to_seq + 1

    if from_row == to_row:
        return {"success": True, "message": "來源與目標席位相同，未做變更"}

    backup_file(ROSTER_XLSX_PATH)
    wb = openpyxl.load_workbook(ROSTER_XLSX_PATH)
    sheet = wb["參賽名單"]

    from_name = sheet.cell(from_row, 3).value
    from_gender = sheet.cell(from_row, 4).value
    from_club = sheet.cell(from_row, 5).value

    to_name = sheet.cell(to_row, 3).value
    to_gender = sheet.cell(to_row, 4).value
    to_club = sheet.cell(to_row, 5).value

    if not from_name and not to_name:
        return {"success": False, "message": "兩個席位皆為空位，未做變更"}

    # 執行調換
    sheet.cell(to_row, 3).value = from_name
    sheet.cell(to_row, 4).value = from_gender
    sheet.cell(to_row, 5).value = from_club

    sheet.cell(from_row, 3).value = to_name
    sheet.cell(from_row, 4).value = to_gender
    sheet.cell(from_row, 5).value = to_club

    wb.save(ROSTER_XLSX_PATH)

    if from_name and to_name:
        msg = f"已成功互換【{from_name}】(第{from_lane}道第{from_seq}位) 與【{to_name}】(第{to_lane}道第{to_seq}位) 的賽道！"
    elif from_name:
        msg = f"已成功將【{from_name}】從第 {from_lane} 道第 {from_seq} 位 移動至 第 {to_lane} 道第 {to_seq} 位！"
    else:
        msg = f"已成功將【{to_name}】從第 {to_lane} 道第 {to_seq} 位 移動至 第 {from_lane} 道第 {from_seq} 位！"

    return {
        "success": True,
        "message": msg,
        "from": {"lane": from_lane, "seq": from_seq, "name": to_name, "gender": to_gender, "club": to_club},
        "to": {"lane": to_lane, "seq": to_seq, "name": from_name, "gender": from_gender, "club": from_club}
    }


def update_player_in_roster(lane, seq, name, gender="男", club=None):
    """
    在「附件四.xlsx」與系統暫存資料庫中手動更新指定球道席位之選手姓名、性別、所屬社
    lane: 1~40, seq: 1~4
    """
    lane = int(lane)
    seq = int(seq)
    if not (1 <= lane <= 40 and 1 <= seq <= 4):
        raise ValueError("球道號碼需介於 1~40，席位序號需介於 1~4")

    row_idx = (lane - 1) * 4 + seq + 1
    backup_file(ROSTER_XLSX_PATH)
    wb = openpyxl.load_workbook(ROSTER_XLSX_PATH)
    sheet = wb["參賽名單"]

    old_name = sheet.cell(row_idx, 3).value
    old_gender = sheet.cell(row_idx, 4).value
    old_club = sheet.cell(row_idx, 5).value

    name = str(name).strip() if name is not None else ""
    gender = str(gender).strip() or "男"

    if not name:
        sheet.cell(row_idx, 3).value = None
        sheet.cell(row_idx, 4).value = None
        if club:
            sheet.cell(row_idx, 5).value = str(club).strip()
            final_club = str(club).strip()
        else:
            sheet.cell(row_idx, 5).value = None
            final_club = None
        message = f"已清空第 {lane} 道第 {seq} 位之選手資料！"
    else:
        sheet.cell(row_idx, 3).value = name
        sheet.cell(row_idx, 4).value = gender
        if club:
            sheet.cell(row_idx, 5).value = str(club).strip()
            final_club = str(club).strip()
        else:
            final_club = old_club
        message = f"已成功將第 {lane} 道第 {seq} 位選手更新為【{name}】({gender})！"

    wb.save(ROSTER_XLSX_PATH)

    # 同步更新 data/roster_data.json 中的選手資料，防止被重新同步覆蓋
    if os.path.exists(ROSTER_DATA_FILE):
        try:
            with open(ROSTER_DATA_FILE, "r", encoding="utf-8") as f:
                roster_data = json.load(f)
            updated = False
            for c_code, c_info in roster_data.items():
                match_club = False
                if final_club and (c_info.get("tag") == final_club or c_code in str(final_club)):
                    match_club = True
                elif old_club and (c_info.get("tag") == old_club or c_code in str(old_club)):
                    match_club = True

                if match_club:
                    for p in list(c_info.get("players", [])):
                        if p.get("name") == old_name or p.get("nickname") == old_name:
                            if name:
                                p["name"] = name
                                p["nickname"] = name
                                p["gender"] = gender
                            else:
                                c_info["players"].remove(p)
                            updated = True
                            break
                elif not updated and old_name:
                    for p in list(c_info.get("players", [])):
                        if p.get("name") == old_name or p.get("nickname") == old_name:
                            if name:
                                p["name"] = name
                                p["nickname"] = name
                                p["gender"] = gender
                            else:
                                c_info["players"].remove(p)
                            updated = True
                            break
            if updated:
                with open(ROSTER_DATA_FILE, "w", encoding="utf-8") as f:
                    json.dump(roster_data, f, ensure_ascii=False, indent=2)
        except Exception as json_err:
            print(f"Warning updating roster_data.json: {json_err}")

    return {
        "success": True,
        "message": message,
        "lane": lane,
        "seq": seq,
        "old_name": old_name,
        "name": name,
        "gender": gender,
        "club": final_club
    }


# ==================== 主流程：郵件檢查與時間回溯同步 ====================

def compute_email_hash(subject, sender, body_text, attachments):
    """計算郵件內容 SHA-256 雜湊，用於比對信件內容是否發生變更"""
    hasher = hashlib.sha256()
    hasher.update((subject or "").encode("utf-8", errors="ignore"))
    hasher.update((sender or "").encode("utf-8", errors="ignore"))
    hasher.update((body_text or "").encode("utf-8", errors="ignore"))
    if attachments:
        for att in sorted(attachments, key=lambda a: a.get("filename", "")):
            hasher.update(att.get("filename", "").encode("utf-8", errors="ignore"))
            hasher.update(att.get("data") or b"")
    return hasher.hexdigest()


def run_sync(dry_run=False, lookback_days=3, extra_buffer_hours=1):
    """
    執行信箱檢查與資料同步：
    - 第一次執行：抓取當下時間往回推算 3 天又 1 小時的 mail 進行偵測（加寬 1 小時防時間差）。
    - 後續執行：自動回溯到「上一次偵測的時間點」再往前多推算 1 小時（防時間差）。
    - 去重與異動修正機制：
        * 已經有抓過且內容無更改的郵件，自動跳過不重複計入。
        * 除非內容有更改，才會重新解析並修正相關紀錄！
    """
    now = datetime.now()
    state = get_sync_state()

    log_messages = []
    def log(msg):
        print(msg)
        log_messages.append(msg)

    log("=" * 60)
    log("  國際扶輪 3523 地區 2026-27 保齡球比賽 AI 自動統計系統")
    log(f"  執行時間: {now.strftime('%Y-%m-%d %H:%M:%S')}")

    # 計算本次查詢的起始時間 (比原定義再多回溯 1 小時防時間差)
    last_sync_iso = state.get("last_sync_time")
    if not last_sync_iso:
        # 初次執行：往回推算 3 天 + 1 小時
        is_first = True
        window_start = now - timedelta(days=lookback_days, hours=extra_buffer_hours)
        log(f"  【初次同步】查詢期間：當下時間往回推算 {lookback_days} 天又 {extra_buffer_hours} 小時防時間差")
    else:
        # 後續執行：回溯到上一次偵測點，再多往前推算 1 小時
        is_first = False
        try:
            last_dt = datetime.fromisoformat(last_sync_iso)
            window_start = last_dt - timedelta(hours=extra_buffer_hours)
            log(f"  【接續同步】查詢期間：自上次偵測點 ({last_dt.strftime('%Y-%m-%d %H:%M:%S')}) 往前加寬 {extra_buffer_hours} 小時防時間差，起始點為 {window_start.strftime('%Y-%m-%d %H:%M:%S')} 至今")
        except Exception:
            window_start = now - timedelta(days=lookback_days, hours=extra_buffer_hours)
            log(f"  【接續同步】解析上次時間異常，預設回溯 {lookback_days} 天又 {extra_buffer_hours} 小時")

    window_end = now
    log(f"  查詢時間區間: {window_start.strftime('%Y-%m-%d %H:%M:%S')} 至 {window_end.strftime('%Y-%m-%d %H:%M:%S')}")
    log("=" * 60)

    clubs = load_clubs_db()

    processed_records = {}
    if os.path.exists(PROCESSED_FILE):
        try:
            with open(PROCESSED_FILE, "r", encoding="utf-8") as f:
                raw_p = json.load(f)
                if isinstance(raw_p, dict):
                    processed_records = raw_p
                elif isinstance(raw_p, list):
                    for item in raw_p:
                        processed_records[str(item)] = {"content_hash": "", "subject": ""}
        except Exception:
            processed_records = {}

    all_registered_data = {}
    if os.path.exists(ROSTER_DATA_FILE):
        with open(ROSTER_DATA_FILE, "r", encoding="utf-8") as f:
            all_registered_data = json.load(f)

    pending_payments = []
    if os.path.exists(PENDING_FILE):
        with open(PENDING_FILE, "r", encoding="utf-8") as f:
            pending_payments = json.load(f)

    # 格式化 IMAP SINCE 日期 (英文月份，避免中文 locale 異常)
    months = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
    imap_since_str = f"{window_start.day:02d}-{months[window_start.month - 1]}-{window_start.year}"

    log(f"連線郵件伺服器 {IMAP_SERVER}...")
    mail = imaplib.IMAP4_SSL(IMAP_SERVER, IMAP_PORT)
    mail.login(MAIL_USER, MAIL_PASS)
    mail.select("INBOX")

    status, data = mail.search(None, f"(SINCE {imap_since_str})")
    mail_ids = data[0].split()
    log(f"該區間在收件匣共有 {len(mail_ids)} 封信件，開始逐一分析新郵件與異動...")

    new_processed = 0
    updated_processed = 0
    skipped_count = 0
    roster_updated = False

    for mid in mail_ids:
        mid_str = mid.decode()

        status, msg_data = mail.fetch(mid, "(RFC822)")
        if not msg_data or not isinstance(msg_data[0], tuple):
            continue

        raw_email = msg_data[0][1]
        msg = email.message_from_bytes(raw_email)
        subject = decode_mime_words(msg.get("Subject", ""))
        sender = decode_mime_words(msg.get("From", ""))
        date_raw = msg.get("Date", "")
        message_id = msg.get("Message-ID", "").strip() or f"{sender}_{date_raw}_{subject}".strip()

        # 解析精確時間，確保在 window_start 之後
        try:
            msg_dt = email.utils.parsedate_to_datetime(date_raw)
            msg_local_dt = msg_dt.astimezone(LOCAL_TZ).replace(tzinfo=None)
            if msg_local_dt < window_start:
                # 早於本次查詢起點，跳過
                continue
        except Exception:
            pass

        body_text, attachments = extract_email_body_and_attachments(msg)
        content_hash = compute_email_hash(subject, sender, body_text, attachments)

        # 檢查是否已抓過 (以唯一且持久的 Message-ID 比對)
        prev_rec = processed_records.get(message_id)
        is_already_seen = prev_rec is not None
        is_content_changed = False

        if is_already_seen:
            prev_hash = prev_rec.get("content_hash", "")
            if prev_hash and prev_hash == content_hash:
                # 已經有抓過且內容無更改 -> 不再抓取
                skipped_count += 1
                continue
            else:
                # 內容有更改 -> 進行修正
                is_content_changed = True
                log(f"\n[偵測到信件內容異動]「{subject}」內容發生變更，進行重新解析與修正！")

        # 1. 檢測上海商銀入帳通知
        bank_info = parse_scsb_deposit_email(subject, body_text, clubs)
        if bank_info:
            action_tag = "【修正入帳】" if is_content_changed else "【新入帳通知】"
            log(f"\n[偵測到上海商銀入帳通知 - {action_tag}]")
            log(f"  主旨: {subject}")
            amt_display = f"NT$ {bank_info['amount']:,}" if bank_info.get("amount") else "（銀行以加密附件提供）"
            log(f"  日期: {bank_info['date']} | 金額: {amt_display}")
            if bank_info.get("memo") or bank_info.get("account_last5"):
                log(f"  備註/摘要: {bank_info['memo']} | 末碼: {bank_info['account_last5']}")

            matched_club = bank_info["club"]
            if bank_info.get("amount"):
                if matched_club:
                    log(f"  --> 成功匹配扶輪社: [{matched_club['code']}] {matched_club['name']}")
                    if not dry_run:
                        note = f"上海商銀入帳"
                        if bank_info["account_last5"]:
                            note += f" (末碼:{bank_info['account_last5']})"
                        update_registration_stats_excel(
                            club=matched_club,
                            remit_date=bank_info["date"],
                            remit_amt=bank_info["amount"],
                            note=note
                        )
                else:
                    matched = False
                    for p in pending_payments:
                        if p.get("amount") == bank_info["amount"] or (p.get("last5") and p.get("last5") == bank_info["account_last5"]):
                            target_club = next((c for c in clubs if c["code"] == p["club_code"]), None)
                            if target_club:
                                log(f"  --> 透過執秘先前報名資料匹配到: [{target_club['code']}] {target_club['name']}")
                                if not dry_run:
                                    update_registration_stats_excel(
                                        club=target_club,
                                        remit_date=bank_info["date"],
                                        remit_amt=bank_info["amount"],
                                        note=f"上海商銀入帳核銷"
                                    )
                                matched = True
                                pending_payments.remove(p)
                                break
                    if not matched:
                        log(f"  [待人工核對] 此筆款項未能直接比對到社名，已記錄至暫存區。")
                        pending_payments.append(bank_info)

            rec_entry = {
                "message_id": message_id,
                "mid": mid_str,
                "subject": subject,
                "sender": sender,
                "date": date_raw,
                "content_hash": content_hash,
                "type": "bank",
                "club_code": matched_club["code"] if matched_club else None,
                "processed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            }
            processed_records[message_id] = rec_entry

            if is_content_changed:
                updated_processed += 1
            else:
                new_processed += 1
            continue

        # 2. 檢測各社執秘保齡球報名信件
        if is_bowling_registration_email(subject, attachments, body_text):
            target_club = None
            parsed_data = None
            for att in attachments:
                fname = att["filename"].lower()
                fdata = att["data"]
                if fname.endswith(".docx"):
                    parsed_data = parse_docx_bytes(fdata)
                    if parsed_data and parsed_data.get("players"):
                        break
                elif fname.endswith(".doc"):
                    parsed_data = parse_doc_bytes(fdata)
                    if parsed_data and parsed_data.get("players"):
                        break
                elif fname.endswith(".pdf"):
                    parsed_data = parse_pdf_bytes(fdata)
                    if parsed_data and parsed_data.get("players"):
                        break

            if not parsed_data or not parsed_data.get("players"):
                parsed_data = parse_text_registration(body_text)

            if parsed_data and (parsed_data.get("players") or parsed_data.get("lunch_only")):
                # 優先從報名表提取社名，其次為信件主旨與寄件者（過濾主辦社關鍵字以防誤判）
                if parsed_data.get("club_name"):
                    target_club = match_club_from_text(parsed_data["club_name"], clubs)
                if not target_club:
                    target_club = match_club_from_text(subject + " " + sender, clubs)
                if not target_club:
                    target_club = match_club_from_text(body_text, clubs)

                if target_club:
                    players = parsed_data.get("players", [])
                    lunch_only = parsed_data.get("lunch_only", [])
                    sponsor = parsed_data.get("sponsor")

                    # 比對該社目前的紀錄是否完全相同
                    curr_reg = all_registered_data.get(target_club["code"])
                    is_club_data_changed = False
                    if curr_reg:
                        old_players = curr_reg.get("players", [])
                        old_lunch = curr_reg.get("lunch_only", [])
                        old_sponsor = curr_reg.get("sponsor")

                        same_p = (
                            len(old_players) == len(players) and
                            all(
                                (p1.get("name") == p2.get("name") or p1.get("nickname") == p2.get("nickname")) and
                                p1.get("gender") == p2.get("gender")
                                for p1, p2 in zip(old_players, players)
                            )
                        )
                        same_l = (len(old_lunch) == len(lunch_only))
                        same_s = (old_sponsor == sponsor)

                        if same_p and same_l and same_s:
                            is_club_data_changed = False
                        else:
                            is_club_data_changed = True
                    else:
                        is_club_data_changed = True

                    if curr_reg and not is_club_data_changed and not is_content_changed:
                        log(f"\n[報名資料無異動] 收到 [{target_club['code']}] {target_club['name']} 報名表，名單與人數與現有紀錄完全相同，維持現有席位不重複寫入。")
                        skipped_count += 1
                    else:
                        action_title = "【修正更正報名】" if (curr_reg or is_content_changed) else "【新報名登記】"
                        log(f"\n[偵測到保齡球比賽相關信件 - {action_title}]")
                        log(f"  主旨: {subject}")
                        log(f"  寄件者: {sender}")
                        log(f"  --> 成功匹配社別: [{target_club['code']}] {target_club['name']}")
                        log(f"  參賽人數: {len(players)} 人 | 僅用餐不參賽: {len(lunch_only)} 人 | 贊助: {sponsor or '一般參賽社'}")
                        for idx, p in enumerate(players, 1):
                            pname = p.get("nickname") or p.get("name")
                            log(f"    選手 {idx:2d}: {pname} ({p.get('gender')})")

                        amt_found = re.search(r"(?:匯款|轉帳)[^\d]*([0-9,]+)\s*元", body_text)
                        last5_found = re.search(r"(?:末[四五]碼|末\s*5\s*碼|帳號末碼)[：:\s]*([0-9]{4,5})", body_text)

                        remit_amt = None
                        remit_date = None
                        last5 = None

                        # 自動辨識是否有水單附件 (圖片)
                        slip_data = {}
                        for att in attachments:
                            fname = att["filename"].lower()
                            if any(k in fname for k in ["水單", "匯款", "轉帳", "繳費", "明細", "收據", "receipt"]) and any(fname.endswith(ext) for ext in [".png", ".jpg", ".jpeg"]):
                                slip_data = ocr_image_slip(att["data"])
                                if slip_data.get("amount"):
                                    break

                        if slip_data.get("amount"):
                            remit_amt = slip_data["amount"]
                            remit_date = slip_data.get("date")
                            last5 = slip_data.get("last5")
                            log(f"  [水單自動辨識] 成功自附件水單辨識金額: NT$ {remit_amt:,} | 末碼: {last5 or '無'} | 日期: {remit_date or '今日'}")

                        if not remit_amt and amt_found:
                            try:
                                remit_amt = int(amt_found.group(1).replace(",", ""))
                            except:
                                pass
                        if not last5 and last5_found:
                            last5 = last5_found.group(1)

                        # 若信件附有水單或提及已匯款，但未能從內文直接取得金額，以報名表合計金額為準
                        has_slip = bool(slip_data.get("amount")) or any("水單" in a["filename"] for a in attachments) or any(k in body_text for k in ["水單", "已匯款", "已轉帳", "匯款水單"])
                        if not remit_amt and has_slip and parsed_data.get("total_amount"):
                            remit_amt = parsed_data["total_amount"]
                            log(f"  [報名費金額核銷] 依水單與報名表合計確認金額: NT$ {remit_amt:,}")

                        if remit_amt and not remit_date:
                            remit_date = datetime.now().strftime("%Y/%m/%d")

                        note_parts = []
                        if parsed_data.get("contact"):
                            note_parts.append(f"執秘:{parsed_data['contact']}")
                        if parsed_data.get("phone"):
                            note_parts.append(f"電話:{parsed_data['phone']}")
                        if last5:
                            note_parts.append(f"末碼:{last5}")
                        if has_slip:
                            note_parts.append("(已附水單)")

                        note_str = " ".join(note_parts)

                        if not dry_run:
                            all_registered_data[target_club["code"]] = {
                                "club_name": target_club["name"],
                                "tag": target_club["tag"],
                                "sponsor": sponsor,
                                "players": players,
                                "lunch_only": lunch_only,
                                "contact": parsed_data.get("contact"),
                                "phone": parsed_data.get("phone"),
                                "updated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                            }

                            update_registration_stats_excel(
                                club=target_club,
                                sponsor=sponsor,
                                players_count=len(players),
                                lunch_count=len(lunch_only),
                                remit_date=remit_date,
                                remit_amt=remit_amt,
                                note=note_str if note_str else None
                            )
                            roster_updated = True

                            if remit_amt:
                                log(f"  --> 報名費 NT$ {remit_amt:,} 已同步登記至「報名統計.xlsx」！")
                            elif last5:
                                pending_payments.append({
                                    "club_code": target_club["code"],
                                    "club_name": target_club["name"],
                                    "amount": remit_amt,
                                    "last5": last5
                                })

                        if is_content_changed or (curr_reg and is_club_data_changed):
                            updated_processed += 1
                        else:
                            new_processed += 1
                else:
                    log(f"  [提示] 發現報名表但無法自動辨識社名，請檢視信件: {subject}")

            rec_entry = {
                "message_id": message_id,
                "mid": mid_str,
                "subject": subject,
                "sender": sender,
                "date": date_raw,
                "content_hash": content_hash,
                "type": "registration",
                "club_code": target_club["code"] if target_club else None,
                "processed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            }
            processed_records[message_id] = rec_entry
            continue

        # 3. 其他非相關信件
        rec_entry = {
            "message_id": message_id,
            "mid": mid_str,
            "subject": subject,
            "sender": sender,
            "date": date_raw,
            "content_hash": content_hash,
            "type": "other",
            "processed_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        }
        processed_records[message_id] = rec_entry

    mail.logout()

    if roster_updated and not dry_run:
        update_roster_excel(all_registered_data, clubs)

    # 更新狀態與下一次預定偵測時間 (3小時後)
    next_sync_dt = window_end + timedelta(hours=AUTO_SYNC_INTERVAL_HOURS)
    if not dry_run:
        new_state = {
            "last_sync_time": window_end.isoformat(),
            "next_scheduled_sync": next_sync_dt.isoformat(),
            "interval_hours": AUTO_SYNC_INTERVAL_HOURS,
            "buffer_hours": extra_buffer_hours,
            "last_window_start": window_start.isoformat(),
            "last_window_end": window_end.isoformat(),
            "is_first_sync": False,
            "total_sync_count": state.get("total_sync_count", 0) + 1
        }
        save_sync_state(new_state)

        with open(PROCESSED_FILE, "w", encoding="utf-8") as f:
            json.dump(processed_records, f, ensure_ascii=False, indent=2)
        with open(ROSTER_DATA_FILE, "w", encoding="utf-8") as f:
            json.dump(all_registered_data, f, ensure_ascii=False, indent=2)
        with open(PENDING_FILE, "w", encoding="utf-8") as f:
            json.dump(pending_payments, f, ensure_ascii=False, indent=2)

    log("=" * 60)
    summary_msg = f"  同步完成！新處理: {new_processed} 封 | 異動更正: {updated_processed} 封 | 已處理略過: {skipped_count} 封"
    log(summary_msg)
    log(f"  下次預定自動偵測時間: {next_sync_dt.strftime('%Y-%m-%d %H:%M:%S')} (每 {AUTO_SYNC_INTERVAL_HOURS} 小時)")
    log("=" * 60)

    return {
        "success": True,
        "is_first": is_first,
        "window_start": window_start.strftime("%Y-%m-%d %H:%M:%S"),
        "window_end": window_end.strftime("%Y-%m-%d %H:%M:%S"),
        "next_sync": next_sync_dt.strftime("%Y-%m-%d %H:%M:%S"),
        "new_processed": new_processed,
        "updated_processed": updated_processed,
        "skipped_count": skipped_count,
        "roster_updated": roster_updated,
        "logs": log_messages
    }


def print_summary_report():
    """輸出即時對帳與報名統計報告"""
    if not os.path.exists(STATS_XLSX_PATH):
        return

    wb = openpyxl.load_workbook(STATS_XLSX_PATH, data_only=True)
    sheet = wb["各社報名資料"]

    total_clubs_registered = 0
    total_players = 0
    total_lunch = 0
    total_receivable = 0
    total_received = 0

    registered_rows = []

    for r in range(2, 82):
        zone = sheet.cell(r, 1).value
        code = sheet.cell(r, 2).value
        name = sheet.cell(r, 3).value
        sponsor = sheet.cell(r, 4).value
        players = sheet.cell(r, 7).value or 0
        lunch = sheet.cell(r, 9).value or 0
        receivable = sheet.cell(r, 11).value or 0
        remit_date = sheet.cell(r, 12).value
        received = sheet.cell(r, 13).value or 0
        unpaid = sheet.cell(r, 14).value or 0

        if players > 0 or lunch > 0 or received > 0 or sponsor:
            total_clubs_registered += 1
            total_players += players
            total_lunch += lunch
            total_receivable += receivable
            total_received += received

            status = "尚未匯款"
            if received > 0:
                if receivable > 0 and received >= receivable:
                    status = "已全額繳清"
                else:
                    status = "部分入帳/待核"

            registered_rows.append({
                "code": code,
                "name": name,
                "sponsor": sponsor or "-",
                "players": players,
                "lunch": lunch,
                "receivable": receivable,
                "received": received,
                "remit_date": remit_date or "-",
                "status": status
            })

    print(f"\n【國際扶輪 3523 地區 2026-27 保齡球比賽 - 即時統計儀表板】")
    print(f"  已報名社數: {total_clubs_registered} 社")
    print(f"  參賽選手總數: {total_players} / 160 人 (剩餘名額: {max(0, 160 - total_players)} 人)")
    print(f"  僅用餐不參賽總人數: {total_lunch} 人")
    print(f"  應收總金額: NT$ {int(total_receivable):,} 元")
    print(f"  實收總金額: NT$ {int(total_received):,} 元")
    print(f"  未收差額: NT$ {int(total_receivable - total_received):,} 元")

    if registered_rows:
        print("\n【各社報名明細】")
        print(f" {'編號':<6} {'社名':<16} {'主協辦':<8} {'選手':<5} {'用餐':<5} {'應收':<8} {'實收':<8} {'匯款日':<10} {'狀態'}")
        print("-" * 80)
        for row in registered_rows:
            print(f" {row['code']:<6} {row['name']:<16} {row['sponsor']:<8} {row['players']:<5} {row['lunch']:<5} {int(row['receivable']):<8} {int(row['received']):<8} {str(row['remit_date']):<10} {row['status']}")
    print("-" * 80 + "\n")


def init_season_roster():
    """
    初始化 附件四.xlsx（參賽名單）：
    保留 1~40 球道與序號 1~4 版型，清空舊選手資料，準備迎接本年度報名。
    """
    backup_file(ROSTER_XLSX_PATH)
    wb = openpyxl.load_workbook(ROSTER_XLSX_PATH)
    sheet = wb["參賽名單"]

    print("正在將「附件四.xlsx」清空初始化為 26-27 年度全新參賽名冊...")
    for slot_idx in range(160):
        row_idx = slot_idx + 2
        lane = (slot_idx // 4) + 1
        seq = (slot_idx % 4) + 1

        sheet.cell(row_idx, 1).value = lane
        for c in range(3, 11):
            sheet.cell(row_idx, c).value = None

    wb.save(ROSTER_XLSX_PATH)
    if os.path.exists(ROSTER_DATA_FILE):
        os.remove(ROSTER_DATA_FILE)
    print("「附件四.xlsx」已成功初始化！")


def reset_all_data(reset_mail_history=False):
    """
    正式上線前一鍵清空所有測試資料，恢復為全新初始狀態：
    1. 備份 報名統計.xlsx 與 附件四.xlsx
    2. 清空 報名統計.xlsx 各社的報名人數、主協辦身分、用餐人數、匯款日、收款金額、備註 (保留所有計算公式)
    3. 清空 附件四.xlsx 參賽名冊中的選手姓名、性別、所屬社別 (保留 40 球道 160 序號版型)
    4. 清空 data/roster_data.json 與 data/pending_payments.json
    5. 重設 data/sync_state.json，使下次收信以當下時間往回推算 3 天
    """
    now_str = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_file(STATS_XLSX_PATH)
    backup_file(ROSTER_XLSX_PATH)

    # 1. 清空 報名統計.xlsx
    wb = openpyxl.load_workbook(STATS_XLSX_PATH, data_only=False)
    sheet = wb["各社報名資料"]
    for r in range(2, 82):
        sheet.cell(r, 4).value = None   # 主協辦身分
        sheet.cell(r, 7).value = None   # 參賽人數
        sheet.cell(r, 9).value = None   # 僅用餐不參賽人數
        sheet.cell(r, 12).value = None  # 匯款日期
        sheet.cell(r, 13).value = None  # 收款金額
        sheet.cell(r, 15).value = None  # 備註
    wb.save(STATS_XLSX_PATH)

    # 2. 清空 附件四.xlsx
    init_season_roster()

    # 3. 清空暫存檔案
    if os.path.exists(ROSTER_DATA_FILE):
        with open(ROSTER_DATA_FILE, "w", encoding="utf-8") as f:
            json.dump({}, f)

    if os.path.exists(PENDING_FILE):
        with open(PENDING_FILE, "w", encoding="utf-8") as f:
            json.dump([], f)

    # 4. 重設同步狀態 (下次點擊將以全新時間往回推算3天)
    initial_sync_state = {
        "last_sync_time": None,
        "next_scheduled_sync": None,
        "interval_hours": AUTO_SYNC_INTERVAL_HOURS,
        "is_first_sync": True,
        "total_sync_count": 0
    }
    save_sync_state(initial_sync_state)

    if reset_mail_history and os.path.exists(PROCESSED_FILE):
        with open(PROCESSED_FILE, "w", encoding="utf-8") as f:
            json.dump({}, f)

    print(">> [系統重設] 所有測試資料已全數清空！系統已恢復為正式上線前的乾淨初始狀態。")
    return True


# ==================== 命令列入口 (CLI Entry) ====================

if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser(description="扶輪保齡球郵件對帳與報名統計助手")
    parser.add_argument("--scan", action="store_true", help="立即連線收信並執行同步")
    parser.add_argument("--report", action="store_true", help="顯示目前對帳與報名總表報告")
    parser.add_argument("--init-season", action="store_true", help="初始化附件四參賽名冊 (清空前年度選手名冊)")
    parser.add_argument("--watch", type=int, nargs="?", const=3, help="定時監聽模式 (預設每 3 小時檢查一次)")
    parser.add_argument("--dry-run", action="store_true", help="測試模式 (不寫入 Excel)")

    args = parser.parse_args()

    if args.init_season:
        init_season_roster()
    elif args.report:
        print_summary_report()
    elif args.watch:
        interval_hours = args.watch
        print(f"啟動定時監聽模式，每 {interval_hours} 小時自動檢查一次信箱...")
        while True:
            try:
                run_sync(dry_run=args.dry_run)
            except Exception as e:
                print(f"[錯誤] 同步時發生例外: {e}")
            print(f"等待 {interval_hours} 小時後再次檢查 (按 Ctrl+C 可停止)...")
            time.sleep(interval_hours * 3600)
    else:
        run_sync(dry_run=args.dry_run)
