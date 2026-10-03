#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
扶輪 3523 地區 2026-27 保齡球比賽 - 網頁管理儀表板
(Rotary Bowling Web Assistant with Auto-Scheduler & Pre-rendered Club Options)
"""

import os
import sys
import json
import time
import threading
from collections import OrderedDict
from datetime import datetime, timedelta
from flask import Flask, render_template_string, jsonify, request, send_file
import openpyxl

# 引入核心同步模組
import bowling_assistant as core

from flask_cors import CORS

app = Flask(__name__)
app.config["JSON_AS_ASCII"] = False
CORS(app, resources={r"/*": {"origins": "*"}}, supports_credentials=True)

@app.after_request
def add_cors_headers(response):
    response.headers["Access-Control-Allow-Origin"] = "*"
    response.headers["Access-Control-Allow-Methods"] = "GET, POST, PUT, DELETE, OPTIONS"
    response.headers["Access-Control-Allow-Headers"] = "Content-Type, Authorization, X-Requested-With"
    return response

@app.before_request
def handle_options_preflight():
    if request.method == "OPTIONS":
        response = app.make_default_options_response()
        response.headers["Access-Control-Allow-Origin"] = "*"
        response.headers["Access-Control-Allow-Methods"] = "GET, POST, PUT, DELETE, OPTIONS"
        response.headers["Access-Control-Allow-Headers"] = "Content-Type, Authorization, X-Requested-With"
        return response

# 記錄即時執行紀錄
recent_logs = []
is_syncing = False
scheduler_thread = None

def add_log(msg):
    timestamp = datetime.now().strftime("%H:%M:%S")
    entry = f"[{timestamp}] {msg}"
    recent_logs.append(entry)
    if len(recent_logs) > 300:
        recent_logs.pop(0)


def background_scheduler_worker():
    """
    每 3 小時自動執行一次的後台監聽線程
    若時間到達下一次預定偵測點，自動觸發同步
    """
    global is_syncing
    while True:
        try:
            state = core.get_sync_state()
            next_sync_iso = state.get("next_scheduled_sync")
            now = datetime.now()

            should_run = False
            if next_sync_iso:
                next_sync_dt = datetime.fromisoformat(next_sync_iso)
                if now >= next_sync_dt:
                    should_run = True

            if should_run and not is_syncing:
                add_log("【定時排程觸發】每 3 小時自動偵測週期到達，開始自動檢測新信件...")
                is_syncing = True
                try:
                    result = core.run_sync(dry_run=False)
                    for l in result.get("logs", []):
                        add_log(l)
                    add_log(f"【定時排程完成】處理了 {result.get('new_processed', 0)} 封信件，下次偵測: {result.get('next_sync')}")
                except Exception as ex:
                    add_log(f"【定時排程例外】: {ex}")
                finally:
                    is_syncing = False

        except Exception as e:
            print(f"[排程監聽錯誤]: {e}")

        time.sleep(15)


def start_scheduler():
    global scheduler_thread
    if scheduler_thread is None or not scheduler_thread.is_alive():
        scheduler_thread = threading.Thread(target=background_scheduler_worker, daemon=True)
        scheduler_thread.start()
        print(">> 每 3 小時自動排程監聽線程已在背景啟動！")

start_scheduler()


HTML_TEMPLATE = r"""
<!DOCTYPE html>
<html lang="zh-TW">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>國際扶輪3523地區 2026-27 保齡球比賽 - AI 自動對帳與報名統計系統</title>
    <!-- Tailwind CSS CDN -->
    <script src="https://cdn.tailwindcss.com"></script>
    <link href="https://cdnjs.cloudflare.com/ajax/libs/font-awesome/6.4.0/css/all.min.css" rel="stylesheet">
    <style>
        body { font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif; }
        .tab-active { border-bottom: 3px solid #2563eb; color: #2563eb; font-weight: 600; }
        .lane-card { transition: all 0.2s; }
        .lane-card:hover { transform: translateY(-2px); box-shadow: 0 4px 12px rgba(0,0,0,0.08); }
        select option { color: #0f172a !important; background-color: #ffffff !important; }
        select optgroup { color: #1e40af !important; background-color: #f1f5f9 !important; font-weight: bold; }
        .slot-selected {
            outline: 2px solid #2563eb;
            background-color: #eff6ff !important;
            box-shadow: 0 0 0 4px rgba(37, 99, 235, 0.2);
            animation: pulse-border 1.5s infinite;
        }
        @keyframes pulse-border {
            0%, 100% { box-shadow: 0 0 0 3px rgba(37, 99, 235, 0.3); }
            50% { box-shadow: 0 0 0 6px rgba(37, 99, 235, 0.15); }
        }
        @media print {
            body * { visibility: hidden !important; }
            #rosterPreviewModal, #rosterPreviewModal * { visibility: visible !important; }
            #rosterPreviewModal {
                position: absolute !important;
                left: 0 !important;
                top: 0 !important;
                width: 100% !important;
                padding: 0 !important;
                background: white !important;
            }
            .no-print { display: none !important; }
        }
    </style>
</head>
<body class="bg-slate-50 text-slate-800 min-h-screen">

    <!-- 頂部導覽列 -->
    <header class="bg-white border-b border-slate-200 sticky top-0 z-50 shadow-sm">
        <div class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8 py-3.5 flex flex-wrap justify-between items-center gap-4">
            <div class="flex items-center space-x-3">
                <div class="w-10 h-10 rounded-xl bg-blue-600 flex items-center justify-center text-white shadow-md shadow-blue-200">
                    <i class="fa-solid fa-bowling-ball text-xl"></i>
                </div>
                <div>
                    <h1 class="text-lg font-bold text-slate-900 leading-tight">國際扶輪 3523 地區 2026-27 保齡球聯誼賽</h1>
                    <p class="text-xs text-slate-500">主辦社：台北市松山扶輪社 ｜ AI 郵件對帳與名冊統計系統</p>
                </div>
            </div>

            <div class="flex items-center space-x-2">
                <button onclick="showServerModal()" id="backendStatusBadge" title="點擊檢視後端連線狀態與啟動指引" class="cursor-pointer inline-flex items-center px-2.5 py-1.5 rounded-lg text-xs font-medium bg-slate-100 text-slate-600 border border-slate-200 hover:bg-slate-200 transition">
                    <span id="backendStatusDot" class="w-2 h-2 rounded-full bg-slate-400 mr-1.5"></span>
                    <span id="backendStatusText">檢測中...</span>
                </button>
                <button onclick="triggerSync()" id="syncBtn" class="inline-flex items-center px-4 py-2 bg-blue-600 hover:bg-blue-700 text-white text-sm font-medium rounded-lg shadow transition disabled:opacity-50">
                    <i class="fa-solid fa-rotate mr-2" id="syncIcon"></i> 立即連線收信對帳
                </button>
                <div class="relative inline-block text-left">
                    <button onclick="toggleDropdown()" class="px-3 py-2 bg-white border border-slate-300 hover:bg-slate-50 text-slate-700 text-sm font-medium rounded-lg transition">
                        <i class="fa-solid fa-download mr-1"></i> 下載 Excel <i class="fa-solid fa-chevron-down text-xs ml-1"></i>
                    </button>
                    <div id="downloadDropdown" class="hidden absolute right-0 mt-2 w-48 rounded-md shadow-lg bg-white ring-1 ring-black ring-opacity-5 z-20">
                        <div class="py-1">
                            <a href="#" onclick="downloadStatsFile();return false;" class="block px-4 py-2 text-sm text-slate-700 hover:bg-slate-100"><i class="fa-solid fa-file-excel text-emerald-600 mr-2"></i>報名統計.xlsx</a>
                            <a href="#" onclick="downloadRosterFile();return false;" class="block px-4 py-2 text-sm text-slate-700 hover:bg-slate-100"><i class="fa-solid fa-file-excel text-blue-600 mr-2"></i>附件四 (參賽名冊).xlsx</a>
                        </div>
                    </div>
                </div>
            </div>
        </div>
    </header>

    <main class="max-w-7xl mx-auto px-4 sm:px-6 lg:px-8 py-6 space-y-6">

        <!-- 自動排程與時間回溯資訊列 -->
        <div class="bg-blue-50/80 border border-blue-200 rounded-xl px-4 py-3 flex flex-wrap justify-between items-center gap-3 text-xs">
            <div class="flex items-center space-x-3">
                <span class="inline-flex items-center px-2.5 py-1 rounded-full bg-emerald-100 text-emerald-800 font-medium">
                    <span class="w-1.5 h-1.5 rounded-full bg-emerald-500 mr-1.5 animate-pulse"></span>
                    每 3 小時自動背景偵測中
                </span>
                <span class="text-slate-600">
                    <i class="fa-regular fa-clock mr-1 text-blue-600"></i> 上次偵測點：
                    <strong id="barLastSync" class="text-slate-900 font-mono">尚未執行</strong>
                </span>
                <span class="text-slate-600">
                    <i class="fa-solid fa-hourglass-half mr-1 text-amber-600"></i> 下次自動偵測：
                    <strong id="barNextSync" class="text-slate-900 font-mono">排程監聽中</strong>
                </span>
            </div>
            <div class="text-slate-500 text-[11px] flex items-center">
                <i class="fa-solid fa-circle-info mr-1 text-blue-500"></i>
                <span id="barWindowTip">初次推算 3 天又 1 小時，後續自動回溯至上次偵測點再往前加寬 1 小時防時間差（已處理無變更者跳過，有更動自動修正）</span>
            </div>
        </div>

        <!-- 數據統計卡片 -->
        <div class="grid grid-cols-1 sm:grid-cols-2 lg:grid-cols-5 gap-4">
            <!-- 已報名社數 -->
            <div class="bg-white rounded-xl p-5 border border-slate-200 shadow-sm flex items-center justify-between">
                <div>
                    <div class="text-xs font-medium text-slate-500 uppercase tracking-wider">已報名社數</div>
                    <div class="text-2xl font-bold text-slate-900 mt-1" id="statClubs">0 <span class="text-xs font-normal text-slate-500">/ 80 社</span></div>
                </div>
                <div class="w-11 h-11 rounded-lg bg-blue-50 text-blue-600 flex items-center justify-center text-lg">
                    <i class="fa-solid fa-users"></i>
                </div>
            </div>

            <!-- 參賽選手總數 -->
            <div class="bg-white rounded-xl p-5 border border-slate-200 shadow-sm flex items-center justify-between">
                <div>
                    <div class="text-xs font-medium text-slate-500 uppercase tracking-wider">參賽選手 (上限 160)</div>
                    <div class="text-2xl font-bold text-slate-900 mt-1"><span id="statPlayers">0</span> <span class="text-xs font-normal text-slate-500">人</span> <span class="text-xs font-normal text-slate-500 ml-1.5" id="statRemaining">餘 160 席</span></div>
                </div>
                <div class="w-11 h-11 rounded-lg bg-amber-50 text-amber-600 flex items-center justify-center text-lg">
                    <i class="fa-solid fa-trophy"></i>
                </div>
            </div>

            <!-- 僅用餐不參賽人數 -->
            <div class="bg-white rounded-xl p-5 border border-slate-200 shadow-sm flex items-center justify-between">
                <div>
                    <div class="text-xs font-medium text-slate-500 uppercase tracking-wider">僅用餐不參賽人數</div>
                    <div class="text-2xl font-bold text-slate-900 mt-1" id="statLunch">0 <span class="text-xs font-normal text-slate-500">人</span></div>
                </div>
                <div class="w-11 h-11 rounded-lg bg-purple-50 text-purple-600 flex items-center justify-center text-lg">
                    <i class="fa-solid fa-utensils"></i>
                </div>
            </div>

            <!-- 應收總額 -->
            <div class="bg-white rounded-xl p-5 border border-slate-200 shadow-sm flex items-center justify-between">
                <div>
                    <div class="text-xs font-medium text-slate-500 uppercase tracking-wider">應收總額</div>
                    <div class="text-2xl font-bold text-slate-900 mt-1" id="statReceivable">$0</div>
                </div>
                <div class="w-11 h-11 rounded-lg bg-indigo-50 text-indigo-600 flex items-center justify-center text-lg">
                    <i class="fa-solid fa-receipt"></i>
                </div>
            </div>

            <!-- 實收與未收 -->
            <div class="bg-white rounded-xl p-5 border border-slate-200 shadow-sm flex items-center justify-between">
                <div>
                    <div class="text-xs font-medium text-slate-500 uppercase tracking-wider">實收總金額</div>
                    <div class="text-2xl font-bold text-emerald-600 mt-1" id="statReceived">$0</div>
                    <div class="text-xs text-rose-500 mt-0.5" id="statUnpaid">未收: $0</div>
                </div>
                <div class="w-11 h-11 rounded-lg bg-emerald-50 text-emerald-600 flex items-center justify-center text-lg">
                    <i class="fa-solid fa-wallet"></i>
                </div>
            </div>
        </div>

        <!-- 名額進度條 -->
        <div class="bg-white rounded-xl p-4 border border-slate-200 shadow-sm">
            <div class="flex justify-between text-xs font-medium text-slate-600 mb-1.5">
                <span><i class="fa-solid fa-gauge-high mr-1 text-blue-600"></i> 160 位參賽額度配額進度</span>
                <span id="progressText">0 / 160 (0%)</span>
            </div>
            <div class="w-full bg-slate-100 rounded-full h-3 overflow-hidden">
                <div id="progressBar" class="bg-blue-600 h-3 rounded-full transition-all duration-500" style="width: 0%"></div>
            </div>
        </div>

        <!-- 功能標籤頁分頁切換 -->
        <div class="border-b border-slate-200 bg-white rounded-t-xl px-4 flex space-x-6">
            <button onclick="switchTab('tab-stats')" id="btn-tab-stats" class="py-3 px-2 text-sm font-medium tab-active flex items-center">
                <i class="fa-solid fa-table-list mr-2"></i> 各社報名與對帳總表
            </button>
            <button onclick="switchTab('tab-roster')" id="btn-tab-roster" class="py-3 px-2 text-sm font-medium text-slate-500 hover:text-slate-700 flex items-center">
                <i class="fa-solid fa-chess-board mr-2"></i> 附件四 40球道排道圖
            </button>
            <button onclick="switchTab('tab-test')" id="btn-tab-test" class="py-3 px-2 text-sm font-medium text-slate-500 hover:text-slate-700 flex items-center">
                <i class="fa-solid fa-vial-circle-check mr-2"></i> 模擬測試輸入
            </button>
            <button onclick="switchTab('tab-log')" id="btn-tab-log" class="py-3 px-2 text-sm font-medium text-slate-500 hover:text-slate-700 flex items-center">
                <i class="fa-solid fa-terminal mr-2"></i> 即時同步日誌與排程
            </button>
        </div>

        <!-- TAB 1: 各社報名與對帳總表 -->
        <div id="tab-stats" class="bg-white rounded-b-xl border border-slate-200 shadow-sm p-5 space-y-4">
            <div class="flex flex-wrap justify-between items-center gap-3">
                <div class="flex items-center space-x-2">
                    <span class="text-sm font-medium text-slate-600">篩選分區：</span>
                    <select id="filterZone" onchange="renderStatsTable()" class="border border-slate-300 rounded-lg px-3 py-1.5 text-sm bg-white text-slate-800 focus:ring-2 focus:ring-blue-500">
                        <option value="ALL">全部 13 個分區</option>
                        <option value="1">第 1 分區</option>
                        <option value="2">第 2 分區</option>
                        <option value="3">第 3 分區</option>
                        <option value="4">第 4 分區</option>
                        <option value="5">第 5 分區</option>
                        <option value="6">第 6 分區</option>
                        <option value="7">第 7 分區</option>
                        <option value="8">第 8 分區</option>
                        <option value="9">第 9 分區</option>
                        <option value="10">第 10 分區</option>
                        <option value="11">第 11 分區</option>
                        <option value="12">第 12 分區</option>
                        <option value="13">第 13 分區</option>
                    </select>

                    <span class="text-sm font-medium text-slate-600 ml-2">繳費狀態：</span>
                    <select id="filterStatus" onchange="renderStatsTable()" class="border border-slate-300 rounded-lg px-3 py-1.5 text-sm bg-white text-slate-800 focus:ring-2 focus:ring-blue-500">
                        <option value="ALL">全部狀態</option>
                        <option value="PAID">已繳款 (綠)</option>
                        <option value="UNPAID">未繳款/報名中 (灰)</option>
                        <option value="DIFF">待對帳/差額 (橙)</option>
                    </select>
                </div>

                <div class="w-64">
                    <input type="text" id="searchClub" oninput="renderStatsTable()" placeholder="搜尋社名或編號..." class="w-full border border-slate-300 rounded-lg px-3 py-1.5 text-sm bg-white text-slate-800 focus:ring-2 focus:ring-blue-500">
                </div>
            </div>

            <!-- 表格 -->
            <div class="overflow-x-auto border border-slate-200 rounded-lg">
                <table class="min-w-full divide-y divide-slate-200 text-sm text-left">
                    <thead class="bg-slate-50 text-slate-600 font-semibold">
                        <tr>
                            <th class="px-3 py-2.5">編號</th>
                            <th class="px-3 py-2.5">社名</th>
                            <th class="px-3 py-2.5">主協辦</th>
                            <th class="px-3 py-2.5 text-center">參賽</th>
                            <th class="px-3 py-2.5 text-right">報名費</th>
                            <th class="px-3 py-2.5 text-center whitespace-nowrap" title="僅用餐不參賽人數">僅用餐不參賽</th>
                            <th class="px-3 py-2.5 text-right">餐費</th>
                            <th class="px-3 py-2.5 text-right font-bold text-slate-800">應收合計</th>
                            <th class="px-3 py-2.5">匯款日</th>
                            <th class="px-3 py-2.5 text-right text-emerald-600 font-bold">實收金額</th>
                            <th class="px-3 py-2.5 text-right">未收差額</th>
                            <th class="px-3 py-2.5 text-center">狀態</th>
                            <th class="px-3 py-2.5">備註</th>
                        </tr>
                    </thead>
                    <tbody id="statsTbody" class="divide-y divide-slate-100 text-slate-700">
                    </tbody>
                </table>
            </div>
        </div>

        <!-- TAB 2: 附件四 40 球道排道圖 (含預覽與選手換道功能) -->
        <div id="tab-roster" class="hidden bg-white rounded-b-xl border border-slate-200 shadow-sm p-5 space-y-4">
            <div class="flex flex-wrap justify-between items-center gap-3 border-b border-slate-100 pb-4">
                <div>
                    <div class="flex items-center space-x-2">
                        <h3 class="text-base font-bold text-slate-900">附件四 參賽名冊排道圖 (共 40 道 × 4 人 = 160 位名額)</h3>
                        <span id="rosterOccupiedBadge" class="text-xs px-2.5 py-0.5 rounded-full bg-blue-100 text-blue-700 font-semibold">已排定 0 / 160 席</span>
                    </div>
                    <p class="text-xs text-slate-500 mt-0.5">
                        <i class="fa-solid fa-arrow-pointer text-blue-500 mr-1"></i>
                        <strong>調換賽道功能：</strong>點擊選手即可選取，再點擊任意「空位」移入或「其他選手」互換；亦支援直接滑鼠拖曳，即時寫入 Excel！
                    </p>
                </div>
                <div class="flex flex-wrap items-center gap-2">
                    <!-- 視圖切換器 -->
                    <div class="inline-flex rounded-lg border border-slate-200 bg-slate-100 p-0.5 text-xs font-medium">
                        <button onclick="switchLanesView('cards')" id="btnViewCards" class="px-3 py-1.5 rounded-md bg-white text-blue-700 shadow-sm font-semibold flex items-center transition">
                            <i class="fa-solid fa-table-cells-large mr-1.5"></i> 40 球道卡片
                        </button>
                        <button onclick="switchLanesView('table')" id="btnViewTable" class="px-3 py-1.5 rounded-md text-slate-600 hover:text-slate-900 flex items-center transition">
                            <i class="fa-solid fa-list-ol mr-1.5"></i> 160 席表格清單
                        </button>
                    </div>

                    <!-- 預覽/列印 Modal 按鈕 -->
                    <button onclick="openRosterPreviewModal()" class="px-3 py-1.5 bg-emerald-600 hover:bg-emerald-700 text-white text-xs font-medium rounded-lg shadow-sm transition flex items-center">
                        <i class="fa-solid fa-file-invoice mr-1.5"></i> 預覽名冊總表 / 列印
                    </button>

                    <button onclick="initSeasonRoster()" class="px-3 py-1.5 border border-rose-300 text-rose-600 hover:bg-rose-50 text-xs font-medium rounded-lg transition flex items-center">
                        <i class="fa-solid fa-trash-can mr-1.5"></i> 清空重設
                    </button>
                    <a href="/download/roster" class="px-3 py-1.5 bg-blue-600 text-white hover:bg-blue-700 text-xs font-medium rounded-lg transition flex items-center">
                        <i class="fa-solid fa-download mr-1.5"></i> 下載 Excel
                    </a>
                </div>
            </div>

            <!-- 卡片模式 -->
            <div id="lanesCardsView">
                <div class="grid grid-cols-1 md:grid-cols-2 lg:grid-cols-4 gap-4" id="lanesContainer">
                </div>
            </div>

            <!-- 表格清單預覽模式 -->
            <div id="lanesTableView" class="hidden space-y-3">
                <div class="flex flex-wrap justify-between items-center gap-3 bg-slate-50 p-3 rounded-lg border border-slate-200 text-xs">
                    <div class="flex flex-wrap items-center gap-3">
                        <div class="flex items-center space-x-1.5">
                            <span class="font-medium text-slate-700"><i class="fa-solid fa-filter mr-1 text-slate-400"></i>球道範圍：</span>
                            <select id="tableFilterLaneRange" onchange="renderLanesTable()" class="border border-slate-300 rounded px-2 py-1 bg-white text-slate-800">
                                <option value="ALL">全部 40 球道 (1~40)</option>
                                <option value="1-10">第 1 ~ 10 球道</option>
                                <option value="11-20">第 11 ~ 20 球道</option>
                                <option value="21-30">第 21 ~ 30 球道</option>
                                <option value="31-40">第 31 ~ 40 球道</option>
                            </select>
                        </div>
                        <label class="flex items-center space-x-1.5 cursor-pointer">
                            <input type="checkbox" id="tableHideEmpty" onchange="renderLanesTable()" class="rounded border-slate-300 text-blue-600">
                            <span class="text-slate-600">僅顯示已排定選手 (隱藏空位)</span>
                        </label>
                    </div>
                    <div>
                        <input type="text" id="tableSearchPlayer" oninput="renderLanesTable()" placeholder="搜尋姓名或社名..." class="border border-slate-300 rounded-lg px-3 py-1 text-xs w-48 bg-white text-slate-800">
                    </div>
                </div>
                <div class="overflow-x-auto border border-slate-200 rounded-lg max-h-[650px] overflow-y-auto">
                    <table class="min-w-full divide-y divide-slate-200 text-xs text-left">
                        <thead class="bg-slate-100 text-slate-700 font-semibold sticky top-0 z-10">
                            <tr>
                                <th class="px-3 py-2.5 text-center w-16">球道</th>
                                <th class="px-3 py-2.5 text-center w-14">席位</th>
                                <th class="px-4 py-2.5 font-bold">選手姓名</th>
                                <th class="px-3 py-2.5 text-center w-16">性別</th>
                                <th class="px-4 py-2.5">所屬扶輪社</th>
                                <th class="px-3 py-2.5 text-center w-16 text-slate-400">第一局</th>
                                <th class="px-3 py-2.5 text-center w-16 text-slate-400">第二局</th>
                                <th class="px-3 py-2.5 text-center w-16 text-slate-400">第三局</th>
                                <th class="px-3 py-2.5 text-center w-16 text-slate-400">火雞獎</th>
                                <th class="px-3 py-2.5 text-center w-28">賽道調換操作</th>
                            </tr>
                        </thead>
                        <tbody id="lanesTableBody" class="divide-y divide-slate-100 text-slate-700 bg-white">
                        </tbody>
                    </table>
                </div>
            </div>
        </div>

        <!-- TAB 3: 模擬測試輸入 -->
        <div id="tab-test" class="hidden bg-white rounded-b-xl border border-slate-200 shadow-sm p-6 space-y-6">
            <div class="max-w-4xl">
                <div class="flex justify-between items-center mb-3">
                    <div>
                        <h3 class="text-base font-bold text-slate-900">功能測試實驗室</h3>
                        <p class="text-xs text-slate-500">在信件進入前，您可手動模擬「執秘寄出報名表」或「上海商銀入帳通知」，以立即驗證系統對帳與自動排列附件四的效果。</p>
                    </div>
                </div>

                <!-- 工具操作列：示範帶入、清空表單、一鍵清空全部測試資料 -->
                <div class="flex flex-wrap justify-between items-center gap-2 mb-5 bg-slate-100/80 p-3 rounded-xl border border-slate-200">
                    <div class="flex items-center space-x-2">
                        <button onclick="quickFillDemo()" class="px-3 py-1.5 bg-indigo-50 hover:bg-indigo-100 text-indigo-700 text-xs font-semibold rounded-lg border border-indigo-200 transition shadow-sm flex items-center">
                            <i class="fa-solid fa-wand-magic-sparkles mr-1.5"></i> 一鍵帶入 12-1 龍門社示範
                        </button>
                        <button onclick="clearTestForm()" class="px-3 py-1.5 bg-white hover:bg-slate-50 text-slate-700 text-xs font-medium rounded-lg border border-slate-300 transition shadow-sm flex items-center">
                            <i class="fa-solid fa-eraser mr-1.5"></i> 清空輸入欄位
                        </button>
                    </div>
                    <button onclick="resetAllTestData()" class="px-3.5 py-1.5 bg-rose-600 hover:bg-rose-700 text-white text-xs font-semibold rounded-lg transition shadow-sm flex items-center">
                        <i class="fa-solid fa-broom mr-1.5"></i> 清空全部測試資料 (正式上線準備)
                    </button>
                </div>

                <div class="grid grid-cols-1 md:grid-cols-2 gap-6">
                    <!-- 模擬 1: 報名表測試 -->
                    <div class="border border-slate-200 rounded-xl p-5 bg-slate-50/70 space-y-3.5">
                        <div class="flex items-center justify-between">
                            <h4 class="font-bold text-sm text-slate-900 flex items-center">
                                <i class="fa-solid fa-file-signature text-blue-600 mr-2"></i> 模擬「執秘寄送報名表」
                            </h4>
                            <span class="text-[11px] text-blue-600 bg-blue-50 px-2 py-0.5 rounded font-medium">寫入統計表 + 排入附件四</span>
                        </div>
                        <div>
                            <label class="text-xs font-semibold text-slate-700 block mb-1">
                                選擇扶輪社 <span class="text-rose-500">*</span>
                            </label>
                            <select id="simClub" onchange="syncClubToBank(this.value)" style="color: #0f172a !important; background-color: #ffffff !important;" class="w-full border border-slate-300 rounded-lg p-2 text-sm font-medium text-slate-900 bg-white shadow-sm focus:ring-2 focus:ring-blue-500">
                                {% for zone, z_clubs in zones.items() %}
                                <optgroup label="【第 {{ zone }} 分區】">
                                    {% for c in z_clubs %}
                                    <option value="{{ c.code }}">[{{ c.code }}] {{ c.name }}</option>
                                    {% endfor %}
                                </optgroup>
                                {% endfor %}
                            </select>
                        </div>
                        <div class="grid grid-cols-2 gap-3">
                            <div>
                                <label class="text-xs font-semibold text-slate-700 block mb-1">主協辦身分</label>
                                <select id="simSponsor" style="color: #0f172a !important; background-color: #ffffff !important;" class="w-full border border-slate-300 rounded-lg p-2 text-sm text-slate-900 bg-white">
                                    <option value="">一般參賽社</option>
                                    <option value="協辦社">協辦社 (5,000元)</option>
                                    <option value="主辦社">共同主辦社 (10,000元)</option>
                                    <option value="贊助社">贊助社 (3,000元)</option>
                                </select>
                            </div>
                            <div>
                                <label class="text-xs font-semibold text-slate-700 block mb-1">僅用餐不參賽人數</label>
                                <input type="number" id="simLunchCount" value="1" min="0" style="color: #0f172a !important; background-color: #ffffff !important;" class="w-full border border-slate-300 rounded-lg p-2 text-sm font-medium text-slate-900 bg-white">
                            </div>
                        </div>
                        <div>
                            <div class="flex justify-between items-center mb-1">
                                <label class="text-xs font-semibold text-slate-700">參賽選手名單 (姓名 性別，每行一人)</label>
                                <span class="text-[10px] text-blue-600 font-medium">支援尊稱/英文名 (例: PP Sure 男、P Lin.C夫人 女、Bank 男)</span>
                            </div>
                            <textarea id="simPlayersText" rows="4" style="color: #0f172a !important; background-color: #ffffff !important;" class="w-full border border-slate-300 rounded-lg p-2.5 text-xs font-mono text-slate-900 bg-white" placeholder="PP Sure 男&#10;P Lin.C夫人 女&#10;PP Borker 男&#10;Bank 男">PP Sure 男
P Lin.C夫人 女
PP Borker 男
Bank 男</textarea>
                        </div>
                        <button onclick="submitSimRegistration()" class="w-full py-2.5 bg-blue-600 hover:bg-blue-700 text-white rounded-lg text-sm font-medium shadow transition">
                            <i class="fa-solid fa-paper-plane mr-1.5"></i> 送出模擬報名 (寫入 Excel)
                        </button>
                    </div>

                    <!-- 模擬 2: 上海商銀入帳測試 -->
                    <div class="border border-slate-200 rounded-xl p-5 bg-slate-50/70 space-y-3.5">
                        <div class="flex items-center justify-between">
                            <h4 class="font-bold text-sm text-slate-900 flex items-center">
                                <i class="fa-solid fa-building-columns text-emerald-600 mr-2"></i> 模擬「上海商銀入帳 Mail」
                            </h4>
                            <span class="text-[11px] text-emerald-600 bg-emerald-50 px-2 py-0.5 rounded font-medium">核銷應收款項</span>
                        </div>
                        <div>
                            <label class="text-xs font-semibold text-slate-700 block mb-1">
                                匯款所屬扶輪社 <span class="text-rose-500">*</span>
                            </label>
                            <select id="simBankClub" style="color: #0f172a !important; background-color: #ffffff !important;" class="w-full border border-slate-300 rounded-lg p-2 text-sm font-medium text-slate-900 bg-white shadow-sm focus:ring-2 focus:ring-emerald-500">
                                {% for zone, z_clubs in zones.items() %}
                                <optgroup label="【第 {{ zone }} 分區】">
                                    {% for c in z_clubs %}
                                    <option value="{{ c.code }}">[{{ c.code }}] {{ c.name }}</option>
                                    {% endfor %}
                                </optgroup>
                                {% endfor %}
                            </select>
                        </div>
                        <div>
                            <label class="text-xs font-semibold text-slate-700 block mb-1">入帳金額 (元)</label>
                            <input type="number" id="simBankAmount" value="5800" style="color: #0f172a !important; background-color: #ffffff !important;" class="w-full border border-slate-300 rounded-lg p-2 text-sm font-medium text-slate-900 bg-white">
                            <p class="text-[11px] text-slate-400 mt-0.5">試算: 4人參賽 4,800 + 1人用席 1,000 = 5,800 元</p>
                        </div>
                        <div>
                            <label class="text-xs font-semibold text-slate-700 block mb-1">入帳日期</label>
                            <input type="text" id="simBankDate" value="{{ today_str }}" style="color: #0f172a !important; background-color: #ffffff !important;" class="w-full border border-slate-300 rounded-lg p-2 text-sm font-medium text-slate-900 bg-white">
                        </div>
                        <div>
                            <label class="text-xs font-semibold text-slate-700 block mb-1">轉帳備註 / 末五碼</label>
                            <input type="text" id="simBankMemo" value="跨行轉入 末碼:39270" style="color: #0f172a !important; background-color: #ffffff !important;" class="w-full border border-slate-300 rounded-lg p-2 text-sm font-medium text-slate-900 bg-white">
                        </div>
                        <button onclick="submitSimDeposit()" class="w-full py-2.5 bg-emerald-600 hover:bg-emerald-700 text-white rounded-lg text-sm font-medium shadow transition">
                            <i class="fa-solid fa-circle-check mr-1.5"></i> 送出模擬入帳 (寫入 Excel)
                        </button>
                    </div>
                </div>
            </div>
        </div>

        <!-- TAB 4: 即時同步日誌與排程 -->
        <div id="tab-log" class="hidden bg-white rounded-b-xl border border-slate-200 shadow-sm p-5 space-y-4">
            <div class="flex justify-between items-center">
                <div class="flex items-center space-x-2">
                    <span class="w-2.5 h-2.5 rounded-full bg-emerald-500 animate-pulse"></span>
                    <h3 class="text-sm font-bold text-slate-900">信箱連線狀態：正常連線中</h3>
                    <span class="text-xs text-slate-500 font-mono">(ssrotary@ms41.hinet.net:993 SSL)</span>
                </div>
                <div class="flex items-center space-x-3 text-xs">
                    <span class="text-slate-500" id="syncStatusSummary">排程模式：每 3 小時自動觸發</span>
                    <button onclick="clearLogs()" class="text-slate-500 hover:text-slate-700 underline">清空日誌</button>
                </div>
            </div>
            <div id="logConsole" class="w-full bg-slate-900 text-emerald-400 font-mono text-xs p-4 rounded-xl h-96 overflow-y-auto leading-relaxed border border-slate-800">
                [系統啟動] 國際扶輪 3523 地區保齡球自動化系統已啟動。<br>
                [排程規則] 第一次點擊收信：往回推算 3 天又 1 小時郵件（加寬 1 小時防時間差）。<br>
                [排程規則] 此後每 3 小時自動偵測；每次回溯到上一次偵測時間點再往前多推算 1 小時。<br>
                [去重規則] 已經抓過且內容無更改之郵件自動略過，不重複計入款項或席位。<br>
                [異動更正] 若郵件內容有異動或執秘寄來更正名單，系統將自動比對並即時修正！<br>
            </div>
        </div>

    </main>

    <!-- 賽道調換底部浮動操作列 (當選中選手時浮現) -->
    <div id="laneSwapBanner" class="hidden fixed bottom-6 left-1/2 -translate-x-1/2 bg-slate-900/95 text-white backdrop-blur-md px-6 py-3.5 rounded-2xl shadow-2xl border border-slate-700 flex flex-wrap items-center gap-4 z-50 transition-all max-w-4xl w-[92%]">
        <div class="flex items-center space-x-3">
            <span class="relative flex h-3 w-3">
                <span class="animate-ping absolute inline-flex h-full w-full rounded-full bg-amber-400 opacity-75"></span>
                <span class="relative inline-flex rounded-full h-3 w-3 bg-amber-500"></span>
            </span>
            <div>
                <div class="text-[11px] text-amber-300 font-medium">🎯 正在更換賽道</div>
                <div class="text-sm font-bold text-white flex items-center space-x-2">
                    <span id="bannerPlayerName">選手姓名</span>
                    <span id="bannerPlayerClub" class="text-xs text-blue-300 font-normal">所屬社別</span>
                    <span class="text-xs px-2 py-0.5 bg-slate-800 text-slate-200 border border-slate-700 rounded" id="bannerPlayerPos">第 1 道 第 1 位</span>
                </div>
            </div>
        </div>

        <div class="h-8 w-px bg-slate-700 hidden sm:block"></div>

        <div class="text-xs text-slate-300 hidden md:block">
            👉 滑鼠點擊目標球道即可<strong>「立即移入」</strong>，免按 Enter 或確認鍵！
        </div>

        <div class="flex items-center space-x-2 ml-auto">
            <span class="text-xs text-slate-400 hidden lg:inline">直接切換球道:</span>
            <select id="quickTargetLane" onchange="executeQuickMove()" title="選取目標球道即刻移入" class="bg-slate-800 text-white text-xs border border-slate-600 rounded px-2 py-1.5 focus:ring-1 focus:ring-emerald-500 cursor-pointer">
            </select>
            <select id="quickTargetSeq" onchange="executeQuickMove()" title="選取目標席位即刻生效" class="bg-slate-800 text-white text-xs border border-slate-600 rounded px-2 py-1.5 focus:ring-1 focus:ring-emerald-500 cursor-pointer">
                <option value="1">席位 1</option>
                <option value="2">席位 2</option>
                <option value="3">席位 3</option>
                <option value="4">席位 4</option>
            </select>
            <button onclick="executeQuickMove()" class="px-3.5 py-1.5 bg-emerald-600 hover:bg-emerald-500 text-white text-xs font-bold rounded-lg shadow transition flex items-center">
                <i class="fa-solid fa-check mr-1"></i> 立即移道
            </button>
            <button onclick="cancelLaneSelection()" class="px-3 py-1.5 bg-slate-700 hover:bg-slate-600 text-slate-300 text-xs rounded-lg transition flex items-center">
                <i class="fa-solid fa-xmark mr-1"></i> 取消
            </button>
        </div>
    </div>

    <!-- 附件四 參賽名冊總表預覽視窗 (Modal) -->
    <div id="rosterPreviewModal" class="hidden fixed inset-0 bg-slate-900/60 backdrop-blur-sm z-50 flex items-center justify-center p-4">
        <div id="previewModalContent" class="bg-white rounded-2xl shadow-2xl border border-slate-200 max-w-5xl w-full max-h-[90vh] flex flex-col overflow-hidden">
            <!-- 視窗標題列 -->
            <div class="px-6 py-4 border-b border-slate-200 flex justify-between items-center bg-slate-50 no-print">
                <div class="flex items-center space-x-3">
                    <div class="w-9 h-9 rounded-lg bg-emerald-600 text-white flex items-center justify-center text-base shadow">
                        <i class="fa-solid fa-file-invoice"></i>
                    </div>
                    <div>
                        <h3 class="text-base font-bold text-slate-900">附件四 參賽名冊總表預覽</h3>
                        <p class="text-xs text-slate-500">全區 40 球道 160 位選手正式名冊（與 Excel 實時同步）</p>
                    </div>
                </div>
                <div class="flex items-center space-x-2">
                    <button onclick="window.print()" class="px-3 py-1.5 bg-slate-800 hover:bg-slate-900 text-white text-xs font-medium rounded-lg shadow transition flex items-center">
                        <i class="fa-solid fa-print mr-1.5"></i> 列印此總表
                    </button>
                    <a href="/download/roster" class="px-3 py-1.5 bg-blue-600 hover:bg-blue-700 text-white text-xs font-medium rounded-lg shadow transition flex items-center">
                        <i class="fa-solid fa-download mr-1.5"></i> 下載 Excel
                    </a>
                    <button onclick="closeRosterPreviewModal()" class="w-8 h-8 rounded-lg hover:bg-slate-200 text-slate-400 hover:text-slate-700 flex items-center justify-center transition text-sm">
                        <i class="fa-solid fa-xmark"></i>
                    </button>
                </div>
            </div>

            <!-- 文件正式抬頭 (列印時亦顯示) -->
            <div class="p-6 pb-2 text-center border-b border-slate-100 bg-white">
                <h2 class="text-xl font-bold text-slate-900 tracking-wider">國際扶輪 3523 地區 2026-27 年度 保齡球聯誼賽</h2>
                <h3 class="text-base font-semibold text-slate-700 mt-1">參賽選手排道名冊（附件四）</h3>
                <div class="flex justify-between items-center text-xs text-slate-500 mt-3 px-2">
                    <span>主辦社：台北市松山扶輪社</span>
                    <span id="previewStatsSummary">比賽球道：40 道 ｜ 參賽席位：160 席 ｜ 已排定 0 人</span>
                    <span>賽程：全三局制</span>
                </div>
            </div>

            <!-- 篩選列 (列印時隱藏) -->
            <div class="px-6 py-2.5 bg-slate-50/80 border-b border-slate-200 flex flex-wrap justify-between items-center gap-3 text-xs no-print">
                <div class="flex items-center space-x-3">
                    <span class="font-medium text-slate-700">球道篩選：</span>
                    <select id="previewFilterLaneRange" onchange="renderPreviewModalTable()" class="border border-slate-300 rounded px-2 py-1 bg-white text-slate-800">
                        <option value="ALL">全部 40 球道 (1~40)</option>
                        <option value="1-10">第 1 ~ 10 球道</option>
                        <option value="11-20">第 11 ~ 20 球道</option>
                        <option value="21-30">第 21 ~ 30 球道</option>
                        <option value="31-40">第 31 ~ 40 球道</option>
                    </select>
                    <label class="flex items-center space-x-1.5 cursor-pointer">
                        <input type="checkbox" id="previewHideEmpty" onchange="renderPreviewModalTable()" class="rounded border-slate-300 text-blue-600">
                        <span class="text-slate-600">僅顯示已排定選手 (隱藏空位)</span>
                    </label>
                </div>
                <div>
                    <input type="text" id="previewSearchPlayer" oninput="renderPreviewModalTable()" placeholder="搜尋姓名或社名..." class="border border-slate-300 rounded-lg px-3 py-1 text-xs w-48 bg-white text-slate-800">
                </div>
            </div>

            <!-- 名冊表格內容 -->
            <div class="p-6 overflow-y-auto flex-1 max-h-[500px]">
                <table class="min-w-full divide-y divide-slate-200 text-xs border border-slate-200 text-left">
                    <thead class="bg-slate-100 text-slate-700 font-semibold sticky top-0">
                        <tr>
                            <th class="px-3 py-2 text-center w-14 border border-slate-200">球道</th>
                            <th class="px-3 py-2 text-center w-12 border border-slate-200">序號</th>
                            <th class="px-4 py-2 border border-slate-200 font-bold">選手姓名</th>
                            <th class="px-3 py-2 text-center w-14 border border-slate-200">性別</th>
                            <th class="px-4 py-2 border border-slate-200">所屬扶輪社</th>
                            <th class="px-3 py-2 text-center w-16 border border-slate-200">第一局</th>
                            <th class="px-3 py-2 text-center w-16 border border-slate-200">第二局</th>
                            <th class="px-3 py-2 text-center w-16 border border-slate-200">第三局</th>
                            <th class="px-3 py-2 text-center w-16 border border-slate-200">火雞獎</th>
                        </tr>
                    </thead>
                    <tbody id="previewModalTableBody" class="divide-y divide-slate-200 bg-white">
                    </tbody>
                </table>
            </div>

            <!-- 底部列 -->
            <div class="px-6 py-3 border-t border-slate-200 bg-slate-50 flex justify-between items-center text-xs text-slate-500 no-print">
                <span>提示：您可以隨時在「40 球道卡片模式」點選選手調整排道，結果將即時反映在此總表與 Excel 檔案中。</span>
                <button onclick="closeRosterPreviewModal()" class="px-4 py-1.5 bg-white border border-slate-300 hover:bg-slate-100 text-slate-700 rounded-lg transition font-medium">
                    關閉
                </button>
            </div>
        </div>
    </div>

    <!-- 本地後端伺服器連線指引 (Modal) -->
    <div id="serverHelpModal" class="hidden fixed inset-0 bg-slate-900/60 backdrop-blur-sm z-50 flex items-center justify-center p-4">
        <div class="bg-white rounded-2xl shadow-2xl border border-slate-200 max-w-lg w-full p-6 space-y-4">
            <div class="flex items-center space-x-3">
                <div class="w-10 h-10 rounded-full bg-blue-100 text-blue-600 flex items-center justify-center text-xl font-bold">
                    <i class="fa-solid fa-server"></i>
                </div>
                <div>
                    <h3 class="text-base font-bold text-slate-900">本地 Python 後端伺服器狀態</h3>
                    <p class="text-xs text-slate-500">HiNet 郵件收發與 Excel 即時排道對帳引擎</p>
                </div>
            </div>

            <div id="serverModalStatusBox" class="p-3.5 rounded-xl border text-xs">
                <!-- 動態注入連線狀態 -->
            </div>

            <div class="space-y-2 text-xs text-slate-600 bg-slate-50 p-3.5 rounded-xl border border-slate-200">
                <div class="font-bold text-slate-800 text-sm mb-1 flex items-center">
                    <i class="fa-solid fa-rocket text-blue-600 mr-1.5"></i> 一鍵啟動後端方式
                </div>
                <p>1. <strong>最簡便方式</strong>：前往保齡球專案資料夾，點兩下執行 <span class="font-mono bg-blue-100 text-blue-800 px-1.5 py-0.5 rounded font-bold">【啟動系統.command】</span> 即可自動啟動後端並開啟瀏覽器！</p>
                <p>2. <strong>終端機執行</strong>：在終端機輸入 <span class="font-mono bg-slate-200 text-slate-800 px-1.5 py-0.5 rounded">.venv/bin/python web_app.py</span></p>
                <p class="text-slate-400 text-[11px] pt-1">（啟動後會常駐於 http://127.0.0.1:5001，支援每 3 小時自動定時收信排程）</p>
            </div>

            <div class="flex justify-between items-center pt-2">
                <button onclick="checkBackendHealth(true)" class="px-4 py-2 bg-blue-600 hover:bg-blue-700 text-white rounded-lg text-xs font-medium transition flex items-center">
                    <i class="fa-solid fa-arrows-rotate mr-1.5"></i> 立即重新檢測連線
                </button>
                <button onclick="closeServerModal()" class="px-4 py-2 bg-white border border-slate-300 hover:bg-slate-100 text-slate-700 rounded-lg text-xs font-medium transition">
                    關閉
                </button>
            </div>
        </div>
    </div>

    <!-- 全域吐司通知 -->
    <div id="globalToast" class="hidden fixed top-5 right-5 z-50 text-white px-4 py-3 rounded-xl shadow-xl flex items-center space-x-2 transition-all duration-300">
        <i id="globalToastIcon" class="fa-solid fa-circle-check text-lg"></i>
        <span id="globalToastMsg" class="text-sm font-medium"></span>
    </div>

    <!-- 拖曳浮動格位 (1:1 選手真實格位，緊貼滑鼠游標，放開左鍵即移入) -->
    <div id="dragGhost"
         class="hidden fixed pointer-events-none z-[9999] bg-white/95 border-2 border-blue-500 rounded-lg py-2 px-2.5 text-xs shadow-2xl flex justify-between items-center select-none ring-4 ring-blue-400/20"
         style="pointer-events: none; will-change: transform, left, top; box-shadow: 0 16px 32px -4px rgba(37,99,235,0.4), 0 8px 16px -2px rgba(0,0,0,0.15); transform: rotate(1.5deg) scale(1.03); transform-origin: center center;">
    </div>

    <!-- JavaScript 互動邏輯 -->
    <script>
        const API_BASE = (window.location.protocol === 'file:' || !window.location.port || window.location.port !== '5001') ? 'http://127.0.0.1:5001' : '';
        let currentData = { stats: [], lanes: [], summary: {}, sync_state: {} };
        let selectedSlot = null;
        let currentLanesViewMode = 'cards';
        let isBackendOnline = false;

        function showServerModal() {
            const m = document.getElementById('serverHelpModal');
            if (m) {
                m.classList.remove('hidden');
                renderServerModalBox();
            }
        }

        function closeServerModal() {
            const m = document.getElementById('serverHelpModal');
            if (m) m.classList.add('hidden');
        }

        function renderServerModalBox() {
            const box = document.getElementById('serverModalStatusBox');
            if (!box) return;
            if (isBackendOnline) {
                box.className = "p-3.5 rounded-xl border bg-emerald-50 border-emerald-200 text-emerald-800 text-xs flex items-center space-x-2";
                box.innerHTML = '<i class="fa-solid fa-circle-check text-base text-emerald-600"></i><div><strong>本地 Python 伺服器運作正常！</strong><br>已成功連線至 http://127.0.0.1:5001，您可以點擊「立即連線收信對帳」進行最新郵件與水單同步。</div>';
            } else {
                box.className = "p-3.5 rounded-xl border bg-rose-50 border-rose-200 text-rose-800 text-xs flex items-center space-x-2";
                box.innerHTML = '<i class="fa-solid fa-triangle-exclamation text-base text-rose-600"></i><div><strong>目前尚未偵測到本地後端伺服器 (http://127.0.0.1:5001)！</strong><br>收信對帳需要透過本機 Python 連線 HiNet 信箱，請依下方步驟啟動服務。</div>';
            }
        }

        function updateBackendStatus(online) {
            isBackendOnline = !!online;
            const badge = document.getElementById('backendStatusBadge');
            const dot = document.getElementById('backendStatusDot');
            const txt = document.getElementById('backendStatusText');
            if (badge && dot && txt) {
                if (online) {
                    badge.className = "cursor-pointer inline-flex items-center px-2.5 py-1.5 rounded-lg text-xs font-medium bg-emerald-50 text-emerald-700 border border-emerald-200 hover:bg-emerald-100 transition";
                    dot.className = "w-2 h-2 rounded-full bg-emerald-500 mr-1.5";
                    txt.innerText = "後端在線 (5001)";
                } else {
                    badge.className = "cursor-pointer inline-flex items-center px-2.5 py-1.5 rounded-lg text-xs font-medium bg-rose-50 text-rose-700 border border-rose-200 hover:bg-rose-100 transition animate-pulse";
                    dot.className = "w-2 h-2 rounded-full bg-rose-500 mr-1.5";
                    txt.innerText = "後端未啟動 (點此啟動)";
                }
            }
            renderServerModalBox();
        }

        async function checkBackendHealth(userTriggered = false) {
            try {
                const res = await fetch(`${API_BASE}/api/health?_t=${Date.now()}`, { method: 'GET', cache: 'no-store' });
                if (res.ok) {
                    updateBackendStatus(true);
                    if (userTriggered) showToast("已成功連線至本地後端伺服器！", "success");
                    return true;
                }
            } catch (e) {}
            updateBackendStatus(false);
            if (userTriggered) showToast("仍無法連線至本地後端 (http://127.0.0.1:5001)，請確認是否已啟動！", "error");
            return false;
        }

        function downloadStatsFile() {
            if (isBackendOnline) {
                window.location.href = `${API_BASE}/download/stats`;
            } else {
                downloadStatsCSV();
            }
        }

        function downloadRosterFile() {
            if (isBackendOnline) {
                window.location.href = `${API_BASE}/download/roster`;
            } else {
                downloadRosterCSV();
            }
        }

        function switchTab(tabId) {
            ['tab-stats', 'tab-roster', 'tab-test', 'tab-log'].forEach(id => {
                document.getElementById(id).classList.add('hidden');
                document.getElementById('btn-' + id).classList.remove('tab-active', 'text-blue-600');
                document.getElementById('btn-' + id).classList.add('text-slate-500');
            });
            document.getElementById(tabId).classList.remove('hidden');
            document.getElementById('btn-' + tabId).classList.add('tab-active', 'text-blue-600');
            document.getElementById('btn-' + tabId).classList.remove('text-slate-500');

            if (tabId === 'tab-roster' && currentData && currentData.lanes && currentData.lanes.length > 0) {
                renderLanes();
                if (currentLanesViewMode === 'table') {
                    renderLanesTable();
                }
            }
        }

        function toggleDropdown() {
            document.getElementById('downloadDropdown').classList.toggle('hidden');
        }

        window.onclick = function(event) {
            if (!event.target.closest('button')) {
                document.getElementById('downloadDropdown').classList.add('hidden');
            }
        }

        function logToConsole(msg) {
            const con = document.getElementById('logConsole');
            const time = new Date().toLocaleTimeString();
            con.innerHTML += `[${time}] ${msg}<br>`;
            con.scrollTop = con.scrollHeight;
        }

        function clearLogs() {
            document.getElementById('logConsole').innerHTML = "";
        }

        function syncClubToBank(val) {
            document.getElementById('simBankClub').value = val;
        }

        function quickFillDemo() {
            document.getElementById('simClub').value = '12-1';
            document.getElementById('simBankClub').value = '12-1';
            document.getElementById('simSponsor').value = '協辦社';
            document.getElementById('simLunchCount').value = '1';
            document.getElementById('simPlayersText').value = `PP Sure 男
P Lin.C夫人 女
PP Borker 男
Bank 男`;
            document.getElementById('simBankAmount').value = '10800'; // 4*1200 + 1000 + 5000 = 10800
            logToConsole("[示範快速帶入] 已帶入 12-1 龍門社 4 位選手與協辦社 5,000 元數據！");
        }

        function clearTestForm() {
            document.getElementById('simPlayersText').value = '';
            document.getElementById('simSponsor').value = '';
            document.getElementById('simLunchCount').value = '0';
            document.getElementById('simBankAmount').value = '0';
            document.getElementById('simBankMemo').value = '';
            logToConsole("[表單重置] 模擬輸入欄位已清空。");
        }

        async function resetAllTestData() {
            const confirmMsg = `⚠️ 確定要清空所有測試資料，恢復為正式上線初始狀態嗎？

此操作將會：
1. 清空「報名統計.xlsx」所有報名與收款紀錄（完整保留全部試算公式）
2. 清空「附件四.xlsx」參賽名冊（保留 40 球道 160 位排道版型）
3. 清除測試暫存，使正式上線收信時以當下時間往回推算 3 天！

(系統已自動在 backups/ 資料夾建立完整時間戳備份)`;

            if (!confirm(confirmMsg)) return;

            logToConsole("正在執行正式上線資料清空重設作業...");
            try {
                const res = await fetch(`${API_BASE}/api/reset-all`, { method: 'POST' });
                const ret = await res.json();
                logToConsole(ret.message);
                clearTestForm();
                await fetchStats();
                alert("✅ " + ret.message);
                switchTab('tab-stats');
            } catch (e) {
                alert("清空重設失敗: " + e);
            }
        }

        async function fetchStats() {
            try {
                const res = await fetch(`${API_BASE}/api/stats?_t=${Date.now()}`);
                if (res.ok) {
                    const data = await res.json();
                    currentData = data;
                    updateSummaryCards(data.summary);
                    updateSyncBar(data.sync_state);
                    renderStatsTable();
                    renderLanes();
                    if (currentLanesViewMode === 'table') {
                        renderLanesTable();
                    }
                    updateBackendStatus(true);
                }
            } catch (err) {
                updateBackendStatus(false);
                console.error("載入失敗:", err);
            }
        }

        function updateSyncBar(state) {
            if (!state) return;
            const barLast = document.getElementById('barLastSync');
            const barNext = document.getElementById('barNextSync');
            const barTip = document.getElementById('barWindowTip');

            if (state.last_sync_time) {
                const dt = new Date(state.last_sync_time);
                barLast.innerText = dt.toLocaleString();
            } else {
                barLast.innerText = "尚未執行 (點擊後抓取前3天)";
            }

            if (state.next_scheduled_sync) {
                const ndt = new Date(state.next_scheduled_sync);
                barNext.innerText = ndt.toLocaleString();
            } else {
                barNext.innerText = "排程待命中";
            }

            if (state.is_first_sync) {
                barTip.innerText = "初次點擊：抓取當下時間往回推算 3 天又 1 小時的 Mail（防時間差）";
            } else {
                barTip.innerText = `本次查詢區間：自上次偵測點再多回溯 1 小時 (${barLast.innerText} - 1h) 至今（已處理無變更者跳過，有更動自動修正）`;
            }
        }

        function updateSummaryCards(s) {
            if (!s) return;
            const elClubs = document.getElementById('statClubs');
            if (elClubs) elClubs.innerHTML = `${s.clubs_count} <span class="text-xs font-normal text-slate-500">/ 80 社</span>`;

            const elPlayers = document.getElementById('statPlayers');
            if (elPlayers) elPlayers.innerText = s.players_count;

            const elRemaining = document.getElementById('statRemaining');
            if (elRemaining) elRemaining.innerText = `餘 ${Math.max(0, 160 - s.players_count)} 席`;

            const elLunch = document.getElementById('statLunch');
            if (elLunch) elLunch.innerHTML = `${s.lunch_count} <span class="text-xs font-normal text-slate-500">人</span>`;

            const elRec = document.getElementById('statReceivable');
            if (elRec) elRec.innerText = '$' + (s.receivable || 0).toLocaleString();

            const elReceived = document.getElementById('statReceived');
            if (elReceived) elReceived.innerText = '$' + (s.received || 0).toLocaleString();

            const elUnpaid = document.getElementById('statUnpaid');
            if (elUnpaid) elUnpaid.innerText = '未收差額: $' + (s.unpaid || 0).toLocaleString();

            const pct = Math.min(100, Math.round((s.players_count / 160) * 100));
            const elBar = document.getElementById('progressBar');
            if (elBar) elBar.style.width = pct + '%';

            const elText = document.getElementById('progressText');
            if (elText) elText.innerText = `${s.players_count} / 160 (${pct}%)`;
        }

        function renderStatsTable() {
            const tbody = document.getElementById('statsTbody');
            const zoneFilter = document.getElementById('filterZone').value;
            const statusFilter = document.getElementById('filterStatus').value;
            const search = document.getElementById('searchClub').value.trim().toLowerCase();

            let rowsHtml = '';
            currentData.stats.forEach(r => {
                if (zoneFilter !== 'ALL' && String(r.zone) !== zoneFilter) return;
                if (search && !r.name.toLowerCase().includes(search) && !r.code.toLowerCase().includes(search)) return;

                const isPaid = r.received >= r.receivable && r.receivable > 0;
                const isUnpaid = r.received === 0;
                if (statusFilter === 'PAID' && !isPaid) return;
                if (statusFilter === 'UNPAID' && (isPaid || r.received > 0)) return;
                if (statusFilter === 'DIFF' && (isPaid || isUnpaid)) return;

                let statusBadge = '<span class="px-2 py-0.5 rounded text-xs bg-slate-100 text-slate-500">未繳</span>';
                if (isPaid) {
                    statusBadge = '<span class="px-2 py-0.5 rounded text-xs bg-emerald-100 text-emerald-700 font-medium">全額已付</span>';
                } else if (r.received > 0) {
                    statusBadge = '<span class="px-2 py-0.5 rounded text-xs bg-amber-100 text-amber-700 font-medium">部分/差額</span>';
                }

                rowsHtml += `
                    <tr class="hover:bg-slate-50/80 transition ${r.players > 0 || r.received > 0 ? 'bg-blue-50/30' : ''}">
                        <td class="px-3 py-2 font-mono text-xs text-slate-500">${r.code}</td>
                        <td class="px-3 py-2 font-medium text-slate-900">${r.name}</td>
                        <td class="px-3 py-2 text-xs text-slate-600">${r.sponsor || '-'}</td>
                        <td class="px-3 py-2 text-center font-semibold ${r.players > 0 ? 'text-blue-600' : 'text-slate-400'}">${r.players || 0}</td>
                        <td class="px-3 py-2 text-right text-xs text-slate-600">${r.player_fee ? '$' + r.player_fee.toLocaleString() : '-'}</td>
                        <td class="px-3 py-2 text-center text-xs text-slate-600">${r.lunch || 0}</td>
                        <td class="px-3 py-2 text-right text-xs text-slate-600">${r.lunch_fee ? '$' + r.lunch_fee.toLocaleString() : '-'}</td>
                        <td class="px-3 py-2 text-right font-bold text-slate-800">${r.receivable ? '$' + r.receivable.toLocaleString() : '$0'}</td>
                        <td class="px-3 py-2 text-xs text-slate-500 font-mono">${r.remit_date || '-'}</td>
                        <td class="px-3 py-2 text-right font-bold text-emerald-600">${r.received ? '$' + r.received.toLocaleString() : '$0'}</td>
                        <td class="px-3 py-2 text-right text-xs ${r.unpaid > 0 ? 'text-rose-500 font-semibold' : 'text-slate-400'}">${r.unpaid ? '$' + r.unpaid.toLocaleString() : '$0'}</td>
                        <td class="px-3 py-2 text-center">${statusBadge}</td>
                        <td class="px-3 py-2 text-xs text-slate-500 max-w-xs truncate" title="${r.note || ''}">${r.note || '-'}</td>
                    </tr>
                `;
            });
            tbody.innerHTML = rowsHtml || '<tr><td colspan="13" class="px-4 py-8 text-center text-slate-400">查無符合條件之社別</td></tr>';
        }

        function showToast(type, msg) {
            const toast = document.getElementById('globalToast');
            const toastMsg = document.getElementById('globalToastMsg');
            const toastIcon = document.getElementById('globalToastIcon');
            if (!toast) return;

            toastMsg.innerText = msg;
            if (type === 'success') {
                toast.className = 'fixed top-5 right-5 z-50 bg-emerald-600 text-white px-5 py-3 rounded-xl shadow-2xl flex items-center space-x-2 transition-all duration-300';
                toastIcon.className = 'fa-solid fa-circle-check text-lg';
            } else {
                toast.className = 'fixed top-5 right-5 z-50 bg-rose-600 text-white px-5 py-3 rounded-xl shadow-2xl flex items-center space-x-2 transition-all duration-300';
                toastIcon.className = 'fa-solid fa-circle-exclamation text-lg';
            }

            clearTimeout(window.toastTimer);
            window.toastTimer = setTimeout(() => {
                toast.className = 'hidden';
            }, 4000);
        }

        function switchLanesView(mode) {
            currentLanesViewMode = mode;
            const cardsView = document.getElementById('lanesCardsView');
            const tableView = document.getElementById('lanesTableView');
            const btnCards = document.getElementById('btnViewCards');
            const btnTable = document.getElementById('btnViewTable');

            if (mode === 'cards') {
                cardsView.classList.remove('hidden');
                tableView.classList.add('hidden');
                btnCards.className = 'px-3 py-1.5 rounded-md bg-white text-blue-700 shadow-sm font-semibold flex items-center transition';
                btnTable.className = 'px-3 py-1.5 rounded-md text-slate-600 hover:text-slate-900 flex items-center transition';
                renderLanes();
            } else {
                cardsView.classList.add('hidden');
                tableView.classList.remove('hidden');
                btnCards.className = 'px-3 py-1.5 rounded-md text-slate-600 hover:text-slate-900 flex items-center transition';
                btnTable.className = 'px-3 py-1.5 rounded-md bg-white text-blue-700 shadow-sm font-semibold flex items-center transition';
                renderLanesTable();
            }
        }

        function renderLanes() {
            const container = document.getElementById('lanesContainer');
            if (!container) return;

            // 確保快速指定球道選單已填入 1~40
            const qLaneSelect = document.getElementById('quickTargetLane');
            if (qLaneSelect && qLaneSelect.options.length === 0) {
                for (let i = 1; i <= 40; i++) {
                    const opt = document.createElement('option');
                    opt.value = i;
                    opt.innerText = `第 ${i} 球道 (點選即移入)`;
                    qLaneSelect.appendChild(opt);
                }
            }

            let cardsHtml = '';
            let occupiedTotal = 0;

            for (let lane = 1; lane <= 40; lane++) {
                const slots = currentData.lanes.filter(item => item.lane === lane);
                const lanePlayerCount = slots.filter(s => s.name).length;
                occupiedTotal += lanePlayerCount;
                const hasPlayers = lanePlayerCount > 0;
                const isOriginLane = selectedSlot && selectedSlot.lane === lane;
                const isTargetLane = selectedSlot && selectedSlot.lane !== lane;

                let slotsHtml = '';
                for (let seq = 1; seq <= 4; seq++) {
                    const slot = slots.find(s => s.seq === seq) || { name: null };
                    const isSelected = selectedSlot && selectedSlot.lane === lane && selectedSlot.seq === seq;

                    if (slot.name) {
                        const genderBadge = slot.gender === '女'
                            ? '<span class="text-[10px] px-1 py-0.2 bg-pink-100 text-pink-600 rounded">女</span>'
                            : '<span class="text-[10px] px-1 py-0.2 bg-blue-100 text-blue-600 rounded">男</span>';

                        let slotClasses = "player-slot flex justify-between items-center py-2 px-2.5 rounded-lg border text-xs transition select-none cursor-grab active:cursor-grabbing group relative ";
                        let titleTip = "按住滑鼠左鍵拖移，放開即完成換道；亦可直接點擊";

                        if (isSelected) {
                            slotClasses += "slot-selected bg-blue-100/90 border-blue-500 ring-2 ring-blue-500 ";
                            titleTip = "目前已選取此選手，點擊其他席位調換，或再次點擊取消";
                        } else if (selectedSlot) {
                            slotClasses += "bg-blue-50/60 border-blue-200 hover:ring-2 hover:ring-amber-400 hover:bg-amber-50 ";
                            titleTip = `點擊立即將【${selectedSlot.name}】與【${slot.name}】互換賽道`;
                        } else {
                            slotClasses += "bg-blue-50/60 border-blue-100 hover:bg-blue-100 hover:border-blue-300 ";
                        }

                        slotsHtml += `
                            <div class="${slotClasses}"
                                 data-lane="${lane}"
                                 data-seq="${seq}"
                                 title="${titleTip}"
                                 draggable="false"
                                 ondragstart="return false;"
                                 onmousedown="startPointerDrag(event, ${lane}, ${seq}, '${slot.name.replace(/'/g, "\\'")}', '${slot.gender || '男'}', '${(slot.club || '').replace(/'/g, "\\'")}')">
                                <div class="flex items-center space-x-1.5 pointer-events-none select-none">
                                    <span class="text-slate-400 font-mono text-[11px] font-semibold">${seq}.</span>
                                    <span class="font-bold text-slate-900">${slot.name}</span>
                                    ${genderBadge}
                                </div>
                                <div class="flex items-center space-x-1 pointer-events-none select-none">
                                    <span class="text-[10px] text-blue-700 truncate max-w-[80px]" title="${slot.club}">${slot.club}</span>
                                    <span class="text-blue-400 text-[10px] opacity-40 group-hover:opacity-100 transition">
                                        <i class="fa-solid fa-up-down-left-right"></i>
                                    </span>
                                </div>
                            </div>
                        `;
                    } else {
                        let emptyClasses = "empty-slot flex items-center justify-between py-2 px-2.5 rounded-lg border border-dashed text-xs transition select-none ";
                        let emptyText = "空位";
                        let titleTip = "空席位";

                        if (selectedSlot) {
                            emptyClasses += "border-emerald-400 bg-emerald-50/70 text-emerald-700 hover:bg-emerald-100 ring-1 ring-emerald-300 cursor-pointer animate-pulse ";
                            emptyText = "👉 點此移入";
                            titleTip = `點擊立即將【${selectedSlot.name}】移動到此空位`;
                        } else {
                            emptyClasses += "border-slate-200 bg-slate-50 text-slate-400 ";
                        }

                        slotsHtml += `
                            <div class="${emptyClasses}"
                                 data-lane="${lane}"
                                 data-seq="${seq}"
                                 title="${titleTip}"
                                 draggable="false"
                                 ondragstart="return false;"
                                 onclick="handleSlotClick(event, ${lane}, ${seq})">
                                <div class="flex items-center pointer-events-none select-none">
                                    <span class="font-mono text-[11px] mr-2 font-semibold text-slate-400">${seq}.</span>
                                    <span class="font-medium">${emptyText}</span>
                                </div>
                                <span class="text-[10px] opacity-60 pointer-events-none select-none">可分配</span>
                            </div>
                        `;
                    }
                }

                let cardClasses = "lane-card bg-white border rounded-xl p-3.5 space-y-2.5 transition-all ";
                let headerRightHtml = `<span class="text-[11px] px-2 py-0.5 rounded-full ${hasPlayers ? 'bg-blue-600 text-white font-medium' : 'bg-slate-100 text-slate-400'}">${lanePlayerCount} / 4 人</span>`;
                let cardAttrs = `data-card-lane="${lane}" `;

                if (isOriginLane) {
                    cardClasses += "border-blue-400 ring-2 ring-blue-300 bg-blue-50/20 ";
                    headerRightHtml = `<span class="text-[11px] px-2 py-0.5 rounded-full bg-blue-600 text-white font-semibold">目前所在道</span>`;
                } else if (isTargetLane) {
                    cardClasses += "border-emerald-400 ring-2 ring-emerald-300/60 bg-emerald-50/20 hover:bg-emerald-50/60 hover:ring-emerald-500 hover:border-emerald-500 cursor-pointer shadow-md ";
                    headerRightHtml = `
                        <button type="button" onclick="event.stopPropagation(); handleCardClick(${lane});" class="text-[11px] px-2.5 py-0.5 rounded bg-emerald-600 hover:bg-emerald-700 text-white font-bold shadow-sm transition flex items-center animate-pulse">
                            <i class="fa-solid fa-arrow-right-to-bracket mr-1"></i> 點此即移入
                        </button>
                    `;
                    cardAttrs += `onclick="handleCardClick(${lane})" title="👉 滑鼠點擊此處或放開左鍵，即刻將【${selectedSlot.name}】移入第 ${lane} 道！" `;
                } else {
                    cardClasses += (hasPlayers ? 'border-blue-200 ring-1 ring-blue-100 ' : 'border-slate-200 ');
                }

                cardsHtml += `
                    <div class="${cardClasses}"
                         ${cardAttrs}>
                        <div class="flex justify-between items-center border-b border-slate-100 pb-2 pointer-events-none">
                            <span class="font-bold text-sm text-slate-900"><i class="fa-solid fa-bowling-ball text-blue-500 mr-1.5"></i> 第 ${lane} 球道</span>
                            ${headerRightHtml}
                        </div>
                        <div class="space-y-1.5">
                            ${slotsHtml}
                        </div>
                    </div>
                `;
            }
            container.innerHTML = cardsHtml;

            const badge = document.getElementById('rosterOccupiedBadge');
            if (badge) {
                badge.innerText = `已排定 ${occupiedTotal} / 160 席`;
            }
        }

        // ==================== 滑鼠按住拖曳・放開左鍵即完成調道 ====================
        let activeDrag = null;
        let dragJustCompleted = false;

        // 全域攔截拖曳結束後引發的點擊事件，避免誤觸選取或取消
        window.addEventListener('click', (e) => {
            if (dragJustCompleted) {
                e.stopPropagation();
                e.preventDefault();
            }
        }, true);

        function startPointerDrag(e, lane, seq, name, gender, club) {
            // 僅響應滑鼠左鍵
            if (e.button !== 0) return;
            e.preventDefault();

            const slotEl = e.currentTarget;
            const rect = slotEl.getBoundingClientRect();

            activeDrag = {
                lane: lane,
                seq: seq,
                name: name,
                gender: gender,
                club: club,
                startX: e.clientX,
                startY: e.clientY,
                // 精準抓取點：滑鼠在格位內部的相對座標（保證滑鼠小手與卡片零距離貼合）
                offsetX: e.clientX - rect.left,
                offsetY: e.clientY - rect.top,
                slotWidth: rect.width,
                slotHeight: rect.height,
                slotEl: slotEl,
                isDragging: false
            };

            window.addEventListener('mousemove', onGlobalMouseMove);
            window.addEventListener('mouseup', onGlobalMouseUp);
        }

        function onGlobalMouseMove(e) {
            if (!activeDrag) return;

            // 滑移超過 3px 認定為拖動調道
            if (!activeDrag.isDragging) {
                if (Math.hypot(e.clientX - activeDrag.startX, e.clientY - activeDrag.startY) > 3) {
                    activeDrag.isDragging = true;
                    if (selectedSlot) {
                        cancelLaneSelection();
                    }

                    // 1:1 精緻複製選手格位，讓滑鼠小手精準按在原本抓取的相對點上
                    const ghost = document.getElementById('dragGhost');
                    if (ghost) {
                        ghost.style.width = activeDrag.slotWidth + 'px';
                        ghost.style.height = activeDrag.slotHeight + 'px';
                        const genderBadge = activeDrag.gender === '女'
                            ? '<span class="text-[10px] px-1 py-0.2 bg-pink-100 text-pink-600 rounded font-medium">女</span>'
                            : '<span class="text-[10px] px-1 py-0.2 bg-blue-100 text-blue-600 rounded font-medium">男</span>';

                        ghost.innerHTML = `
                            <div class="flex items-center space-x-1.5 pointer-events-none select-none">
                                <span class="text-blue-600 font-mono text-[11px] font-bold">${activeDrag.seq}.</span>
                                <span class="font-bold text-slate-900 text-xs">${activeDrag.name}</span>
                                ${genderBadge}
                            </div>
                            <div class="flex items-center space-x-1.5 pointer-events-none select-none">
                                <span class="text-[10px] text-blue-700 font-semibold truncate max-w-[85px]">${activeDrag.club || '未定'}</span>
                                <span class="text-[9px] bg-emerald-600 text-white px-1.5 py-0.5 rounded font-bold shadow animate-pulse">放開即移入</span>
                            </div>
                        `;
                        ghost.classList.remove('hidden');
                        ghost.style.left = (e.clientX - activeDrag.offsetX) + 'px';
                        ghost.style.top = (e.clientY - activeDrag.offsetY) + 'px';
                    }

                    // 原始席位呈現虛線凹槽占位感，視覺精緻度大增
                    if (activeDrag.slotEl) {
                        activeDrag.slotEl.classList.add('opacity-30', 'border-dashed', 'border-blue-400', 'bg-blue-50/40');
                    }

                    document.body.style.userSelect = 'none';
                    document.body.style.cursor = 'grabbing';
                } else {
                    return;
                }
            }

            // 讓格位精準無延遲貼齊滑鼠指針（小手），距離為 0！
            const ghost = document.getElementById('dragGhost');
            if (ghost) {
                ghost.style.left = (e.clientX - activeDrag.offsetX) + 'px';
                ghost.style.top = (e.clientY - activeDrag.offsetY) + 'px';
            }

            // 清理舊高亮
            document.querySelectorAll('.drag-hover-active').forEach(el => {
                el.classList.remove('drag-hover-active', 'ring-4', 'ring-emerald-500/80', 'bg-emerald-50/60', 'border-emerald-500', 'shadow-md');
            });
            document.querySelectorAll('.drag-slot-hover-active').forEach(el => {
                el.classList.remove('drag-slot-hover-active', 'ring-2', 'ring-emerald-500', 'bg-emerald-100', 'border-emerald-500', 'scale-[1.02]');
            });

            // 確保座標限制在視窗內，防護邊緣 null
            const cx = Math.max(1, Math.min(window.innerWidth - 2, e.clientX));
            const cy = Math.max(1, Math.min(window.innerHeight - 2, e.clientY));
            const el = document.elementFromPoint(cx, cy);

            if (el) {
                // 高亮目標球道卡片
                const card = el.closest('[data-card-lane]');
                if (card) {
                    const hoverLane = parseInt(card.getAttribute('data-card-lane'));
                    if (hoverLane !== activeDrag.lane) {
                        card.classList.add('drag-hover-active', 'ring-4', 'ring-emerald-500/80', 'bg-emerald-50/60', 'border-emerald-500', 'shadow-md');
                    }
                }
                // 高亮目標席位 (特定空位或選手)
                const slotEl = el.closest('[data-lane][data-seq]');
                if (slotEl) {
                    const sLane = parseInt(slotEl.getAttribute('data-lane'));
                    const sSeq = parseInt(slotEl.getAttribute('data-seq'));
                    if (sLane !== activeDrag.lane || sSeq !== activeDrag.seq) {
                        slotEl.classList.add('drag-slot-hover-active', 'ring-2', 'ring-emerald-500', 'bg-emerald-100', 'border-emerald-500', 'scale-[1.02]');
                    }
                }
            }
        }

        function onGlobalMouseUp(e) {
            window.removeEventListener('mousemove', onGlobalMouseMove);
            window.removeEventListener('mouseup', onGlobalMouseUp);
            document.body.style.userSelect = '';
            document.body.style.cursor = '';

            const ghost = document.getElementById('dragGhost');
            if (ghost) {
                ghost.classList.add('hidden');
            }

            document.querySelectorAll('.drag-hover-active').forEach(el => {
                el.classList.remove('drag-hover-active', 'ring-4', 'ring-emerald-500/80', 'bg-emerald-50/60', 'border-emerald-500', 'shadow-md');
            });
            document.querySelectorAll('.drag-slot-hover-active').forEach(el => {
                el.classList.remove('drag-slot-hover-active', 'ring-2', 'ring-emerald-500', 'bg-emerald-100', 'border-emerald-500', 'scale-[1.02]');
            });

            if (!activeDrag) return;

            // 恢復原始席位半透明樣式
            if (activeDrag.slotEl) {
                activeDrag.slotEl.classList.remove('opacity-30', 'border-dashed', 'border-blue-400', 'bg-blue-50/40');
            }

            const dragInfo = activeDrag;
            activeDrag = null;

            // ⚡ 關鍵特色：按住滑鼠左鍵滑移到達目標後，「放開滑鼠左鍵即刻完成調道」！免任何多餘點擊！
            if (dragInfo.isDragging) {
                dragJustCompleted = true;
                setTimeout(() => { dragJustCompleted = false; }, 350);

                const cx = Math.max(1, Math.min(window.innerWidth - 2, e.clientX));
                const cy = Math.max(1, Math.min(window.innerHeight - 2, e.clientY));
                const el = document.elementFromPoint(cx, cy);

                if (el) {
                    // 1. 若放開在特定席位上（空位或其他選手）
                    const slotEl = el.closest('[data-lane][data-seq]');
                    if (slotEl) {
                        const toLane = parseInt(slotEl.getAttribute('data-lane'));
                        const toSeq = parseInt(slotEl.getAttribute('data-seq'));
                        if (toLane !== dragInfo.lane || toSeq !== dragInfo.seq) {
                            executeMove(dragInfo.lane, dragInfo.seq, toLane, toSeq);
                            return;
                        }
                    }

                    // 2. 若放開在球道卡片任何區域（自動分配至該球道的第一個空位，滿員則與席位 1 互換）
                    const card = el.closest('[data-card-lane]');
                    if (card) {
                        const toLane = parseInt(card.getAttribute('data-card-lane'));
                        if (toLane !== dragInfo.lane) {
                            const slots = currentData.lanes.filter(item => item.lane === toLane);
                            const emptySlot = slots.find(s => !s.name);
                            const toSeq = emptySlot ? emptySlot.seq : 1;
                            executeMove(dragInfo.lane, dragInfo.seq, toLane, toSeq);
                            return;
                        }
                    }
                }
                return;
            }

            // 若使用者沒有滑動拖曳（只是原地輕點選手）：直接執行選取或取消
            handleSlotClick(null, dragInfo.lane, dragInfo.seq);
        }

        document.addEventListener('keydown', (e) => {
            if (e.key === 'Escape') {
                if (activeDrag && activeDrag.isDragging) {
                    if (activeDrag.slotEl) {
                        activeDrag.slotEl.classList.remove('opacity-30', 'border-dashed', 'border-blue-400', 'bg-blue-50/40');
                    }
                    activeDrag = null;
                    window.removeEventListener('mousemove', onGlobalMouseMove);
                    window.removeEventListener('mouseup', onGlobalMouseUp);
                    document.body.style.userSelect = '';
                    document.body.style.cursor = '';
                    const ghost = document.getElementById('dragGhost');
                    if (ghost) ghost.classList.add('hidden');
                    document.querySelectorAll('.drag-hover-active').forEach(el => {
                        el.classList.remove('drag-hover-active', 'ring-4', 'ring-emerald-500/80', 'bg-emerald-50/60', 'border-emerald-500', 'shadow-md');
                    });
                    document.querySelectorAll('.drag-slot-hover-active').forEach(el => {
                        el.classList.remove('drag-slot-hover-active', 'ring-2', 'ring-emerald-500', 'bg-emerald-100', 'border-emerald-500', 'scale-[1.02]');
                    });
                } else if (selectedSlot) {
                    cancelLaneSelection();
                }
            }
        });

        function handleCardClick(lane) {
            if (dragJustCompleted) return;
            if (!selectedSlot) return;
            if (selectedSlot.lane === lane) {
                cancelLaneSelection();
                return;
            }
            // 尋找目標球道的第一個空席位；若皆滿員則與席位 1 互換
            const slots = currentData.lanes.filter(item => item.lane === lane);
            const emptySlot = slots.find(s => !s.name);
            const targetSeq = emptySlot ? emptySlot.seq : 1;
            executeMove(selectedSlot.lane, selectedSlot.seq, lane, targetSeq);
        }

        function handleSlotClick(event, lane, seq) {
            if (event) event.stopPropagation();
            if (dragJustCompleted) return;

            const slot = currentData.lanes.find(item => item.lane === lane && item.seq === seq) || { name: null };

            if (!selectedSlot) {
                if (!slot.name) {
                    showToast('error', '此席位為空位，請先點選有參賽選手的席位！');
                    return;
                }
                selectSlot(lane, seq, slot.name, slot.gender, slot.club);
                return;
            }

            if (selectedSlot.lane === lane && selectedSlot.seq === seq) {
                cancelLaneSelection();
                return;
            }

            executeMove(selectedSlot.lane, selectedSlot.seq, lane, seq);
        }

        function selectSlot(lane, seq, name, gender, club) {
            selectedSlot = { lane, seq, name, gender, club };
            document.getElementById('bannerPlayerName').innerText = name;
            document.getElementById('bannerPlayerClub').innerText = club || '無社名';
            document.getElementById('bannerPlayerPos').innerText = `第 ${lane} 道 第 ${seq} 位`;
            document.getElementById('quickTargetLane').value = lane;
            document.getElementById('quickTargetSeq').value = seq;
            document.getElementById('laneSwapBanner').classList.remove('hidden');

            renderLanes();
            if (currentLanesViewMode === 'table') {
                renderLanesTable();
            }
        }

        function cancelLaneSelection() {
            selectedSlot = null;
            document.getElementById('laneSwapBanner').classList.add('hidden');
            renderLanes();
            if (currentLanesViewMode === 'table') {
                renderLanesTable();
            }
        }

        function executeQuickMove() {
            if (!selectedSlot) return;
            const targetLane = parseInt(document.getElementById('quickTargetLane').value);
            const targetSeq = parseInt(document.getElementById('quickTargetSeq').value);
            executeMove(selectedSlot.lane, selectedSlot.seq, targetLane, targetSeq);
        }

        async function executeMove(fromLane, fromSeq, toLane, toSeq) {
            if (fromLane === toLane && fromSeq === toSeq) {
                cancelLaneSelection();
                return;
            }

            logToConsole(`[賽道更換] 正在將 第 ${fromLane} 道第 ${fromSeq} 位 移動至 第 ${toLane} 道第 ${toSeq} 位...`);
            cancelLaneSelection();

            try {
                const res = await fetch(`${API_BASE}/api/swap-lane`, {
                    method: 'POST',
                    headers: { 'Content-Type': 'application/json' },
                    body: JSON.stringify({
                        from_lane: fromLane,
                        from_seq: fromSeq,
                        to_lane: toLane,
                        to_seq: toSeq
                    })
                });
                const result = await res.json();
                if (result.success) {
                    logToConsole("✅ " + result.message);
                    showToast('success', result.message);
                    await fetchStats();
                } else {
                    logToConsole("❌ 更換失敗: " + result.message);
                    showToast('error', result.message);
                }
            } catch (err) {
                logToConsole("❌ 請求失敗: " + err);
                showToast('error', "更換請求失敗: " + err);
            }
        }

        function renderLanesTable() {
            const tbody = document.getElementById('lanesTableBody');
            if (!tbody) return;

            const filterRange = document.getElementById('tableFilterLaneRange').value;
            const hideEmpty = document.getElementById('tableHideEmpty').checked;
            const search = document.getElementById('tableSearchPlayer').value.trim().toLowerCase();

            let rowsHtml = '';

            currentData.lanes.forEach(slot => {
                if (filterRange !== 'ALL') {
                    const [minL, maxL] = filterRange.split('-').map(Number);
                    if (slot.lane < minL || slot.lane > maxL) return;
                }
                if (hideEmpty && !slot.name) return;
                if (search) {
                    const sName = (slot.name || '').toLowerCase();
                    const sClub = (slot.club || '').toLowerCase();
                    if (!sName.includes(search) && !sClub.includes(search)) return;
                }

                const isSelected = selectedSlot && selectedSlot.lane === slot.lane && selectedSlot.seq === slot.seq;
                const genderBadge = slot.gender === '女'
                    ? '<span class="px-1.5 py-0.5 bg-pink-100 text-pink-600 rounded text-[11px] font-medium">女</span>'
                    : (slot.gender === '男' ? '<span class="px-1.5 py-0.5 bg-blue-100 text-blue-600 rounded text-[11px] font-medium">男</span>' : '-');

                let actionBtn = '';
                if (slot.name) {
                    if (isSelected) {
                        actionBtn = '<button onclick="cancelLaneSelection()" class="px-2 py-1 bg-amber-100 text-amber-800 rounded font-medium text-xs hover:bg-amber-200 transition">取消選取</button>';
                    } else if (selectedSlot) {
                        actionBtn = `<button onclick="executeMove(${selectedSlot.lane}, ${selectedSlot.seq}, ${slot.lane}, ${slot.seq})" class="px-2 py-1 bg-amber-500 text-white rounded font-medium text-xs hover:bg-amber-600 transition shadow-sm">與此互換</button>`;
                    } else {
                        actionBtn = `<button onclick="handleSlotClick(event, ${slot.lane}, ${slot.seq})" class="px-2 py-1 bg-blue-50 text-blue-600 hover:bg-blue-100 rounded border border-blue-200 text-xs transition"><i class="fa-solid fa-arrow-right-arrow-left mr-1"></i>調換賽道</button>`;
                    }
                } else {
                    if (selectedSlot) {
                        actionBtn = `<button onclick="executeMove(${selectedSlot.lane}, ${selectedSlot.seq}, ${slot.lane}, ${slot.seq})" class="px-2 py-1 bg-emerald-600 text-white rounded font-medium text-xs hover:bg-emerald-700 transition shadow-sm"><i class="fa-solid fa-arrow-down mr-1"></i>移入此位</button>`;
                    } else {
                        actionBtn = '<span class="text-slate-300 text-xs">-</span>';
                    }
                }

                rowsHtml += `
                    <tr class="hover:bg-slate-50 transition ${isSelected ? 'bg-blue-50/80 font-medium' : (slot.name ? '' : 'text-slate-400 bg-slate-50/30')}">
                        <td class="px-3 py-2 text-center font-mono font-semibold text-slate-800">${slot.lane}</td>
                        <td class="px-3 py-2 text-center font-mono text-slate-500">${slot.seq}</td>
                        <td class="px-4 py-2 font-bold ${slot.name ? 'text-slate-900' : 'text-slate-400 font-normal italic'}">${slot.name || '(空位)'}</td>
                        <td class="px-3 py-2 text-center">${genderBadge}</td>
                        <td class="px-4 py-2 text-xs text-slate-700">${slot.club || '-'}</td>
                        <td class="px-3 py-2 text-center text-slate-300">-</td>
                        <td class="px-3 py-2 text-center text-slate-300">-</td>
                        <td class="px-3 py-2 text-center text-slate-300">-</td>
                        <td class="px-3 py-2 text-center text-slate-300">-</td>
                        <td class="px-3 py-2 text-center">${actionBtn}</td>
                    </tr>
                `;
            });

            tbody.innerHTML = rowsHtml || '<tr><td colspan="10" class="py-8 text-center text-slate-400">查無符合條件之席位資料</td></tr>';
        }

        function openRosterPreviewModal() {
            document.getElementById('rosterPreviewModal').classList.remove('hidden');
            renderPreviewModalTable();
        }

        function closeRosterPreviewModal() {
            document.getElementById('rosterPreviewModal').classList.add('hidden');
        }

        function renderPreviewModalTable() {
            const tbody = document.getElementById('previewModalTableBody');
            if (!tbody) return;

            const filterRange = document.getElementById('previewFilterLaneRange').value;
            const hideEmpty = document.getElementById('previewHideEmpty').checked;
            const search = document.getElementById('previewSearchPlayer').value.trim().toLowerCase();

            let rowsHtml = '';
            let occupied = 0;

            currentData.lanes.forEach(slot => {
                if (slot.name) occupied++;

                if (filterRange !== 'ALL') {
                    const [minL, maxL] = filterRange.split('-').map(Number);
                    if (slot.lane < minL || slot.lane > maxL) return;
                }
                if (hideEmpty && !slot.name) return;
                if (search) {
                    const sName = (slot.name || '').toLowerCase();
                    const sClub = (slot.club || '').toLowerCase();
                    if (!sName.includes(search) && !sClub.includes(search)) return;
                }

                const genderStr = slot.gender || (slot.name ? '男' : '-');
                rowsHtml += `
                    <tr class="hover:bg-slate-50 text-xs border-b border-slate-100 ${slot.name ? '' : 'text-slate-300'}">
                        <td class="px-3 py-2 text-center font-mono font-semibold border-r border-slate-100">${slot.lane}</td>
                        <td class="px-3 py-2 text-center font-mono border-r border-slate-100">${slot.seq}</td>
                        <td class="px-4 py-2 font-bold ${slot.name ? 'text-slate-900' : 'text-slate-300 italic'} border-r border-slate-100">${slot.name || '(空位)'}</td>
                        <td class="px-3 py-2 text-center border-r border-slate-100">${genderStr}</td>
                        <td class="px-4 py-2 text-slate-700 border-r border-slate-100">${slot.club || '-'}</td>
                        <td class="px-3 py-2 text-center text-slate-300 border-r border-slate-100"></td>
                        <td class="px-3 py-2 text-center text-slate-300 border-r border-slate-100"></td>
                        <td class="px-3 py-2 text-center text-slate-300 border-r border-slate-100"></td>
                        <td class="px-3 py-2 text-center text-slate-300"></td>
                    </tr>
                `;
            });

            tbody.innerHTML = rowsHtml || '<tr><td colspan="9" class="py-8 text-center text-slate-400">查無符合條件的名冊紀錄</td></tr>';
            const sumEl = document.getElementById('previewStatsSummary');
            if (sumEl) {
                sumEl.innerText = `比賽球道：40 道 ｜ 總席位：160 席 ｜ 已排定 ${occupied} 人 ｜ 剩餘名額 ${Math.max(0, 160 - occupied)} 席`;
            }
        }

        async function triggerSync() {
            const btn = document.getElementById('syncBtn');
            const icon = document.getElementById('syncIcon');
            btn.disabled = true;
            icon.classList.add('fa-spin');
            
            logToConsole("正在連線 ssrotary@ms41.hinet.net 進行收信對帳...");

            try {
                const res = await fetch(`${API_BASE}/api/sync`, { method: 'POST' });
                if (!res.ok) {
                    const errJson = await res.json().catch(() => ({}));
                    throw new Error(errJson.message || `伺服器回應異常 (HTTP ${res.status})`);
                }
                const result = await res.json();
                logToConsole("✅ " + (result.message || "同步完成！"));
                if (result.logs && Array.isArray(result.logs)) {
                    result.logs.forEach(l => logToConsole(l));
                }
                updateBackendStatus(true);
                await fetchStats();
                showToast(result.message || "同步完成！", "success");
            } catch (err) {
                updateBackendStatus(false);
                logToConsole("❌ 同步失敗: " + err.message);
                showToast("連線後端失敗: " + err.message, "error");
                showServerModal();
            } finally {
                btn.disabled = false;
                icon.classList.remove('fa-spin');
            }
        }

        function parseSimPlayerLine(rawLine) {
            let line = rawLine.trim();
            if (!line) return null;

            // 1. 去除開頭序號，例如 "1.", "1 ", "1、", "#1", "(1)"
            line = line.replace(/^\s*(?:#?\d+[\.\、\s\-\:\)]+|\(\d+\))\s*/, '').trim();
            if (!line) return null;

            // 2. 扶輪社參賽選手常包含尊稱、英文名或夫人稱謂
            // 例如: "PP Sure 男", "P Lin.C夫人 女", "PP Borker 男", "Bank 男", "王大明 男", "林小華 女"
            // 以末尾或以空格隔開的「男/女/M/F」作為性別，前面的全部視為完整人名！
            let name = line;
            let gender = "男";

            const genderMatch = line.match(/^(.*?)[,\s\t]+([男女]|[Mm]ale|[Ff]emale|[MFmf])(?:\s+.*)?$/);
            if (genderMatch) {
                name = genderMatch[1].trim();
                const gStr = genderMatch[2].trim().toLowerCase();
                if (gStr === '女' || gStr === 'f' || gStr === 'female') {
                    gender = '女';
                } else {
                    gender = '男';
                }
            } else {
                // 若未標示性別：若含有「夫人」則判定為女，其餘預設為男
                if (name.includes('夫人')) {
                    gender = '女';
                } else {
                    gender = '男';
                }
            }

            // 清理姓名頭尾多餘符號
            name = name.replace(/^[\,\，\s]+|[\,\，\s]+$/g, '').trim();
            if (!name) return null;

            return {
                name: name,
                nickname: name,
                gender: gender
            };
        }

        async function submitSimRegistration() {
            const club_code = document.getElementById('simClub').value;
            const sponsor = document.getElementById('simSponsor').value;
            const lunch_count = parseInt(document.getElementById('simLunchCount').value) || 0;
            const text = document.getElementById('simPlayersText').value.trim();

            const players = [];
            text.split('\\n').forEach(line => {
                const p = parseSimPlayerLine(line);
                if (p) {
                    players.push(p);
                }
            });

            logToConsole(`[模擬測試] 送出 [${club_code}] 報名: ${players.length} 位選手, ${lunch_count} 人僅用餐不參賽...`);
            try {
                const res = await fetch(`${API_BASE}/api/simulate-registration`, {
                    method: 'POST',
                    headers: { 'Content-Type': 'application/json' },
                    body: JSON.stringify({ club_code, sponsor, lunch_count, players })
                });
                const ret = await res.json();
                logToConsole(ret.message);
                await fetchStats();
                switchTab('tab-stats');
            } catch (e) {
                logToConsole("送出失敗: " + e);
            }
        }

        async function submitSimDeposit() {
            const club_code = document.getElementById('simBankClub').value;
            const amount = parseInt(document.getElementById('simBankAmount').value) || 0;
            const date = document.getElementById('simBankDate').value;
            const memo = document.getElementById('simBankMemo').value;

            logToConsole(`[模擬測試] 送出上海商銀入帳: [${club_code}] 金額 $${amount}...`);
            try {
                const res = await fetch(`${API_BASE}/api/simulate-deposit`, {
                    method: 'POST',
                    headers: { 'Content-Type': 'application/json' },
                    body: JSON.stringify({ club_code, amount, date, memo })
                });
                const ret = await res.json();
                logToConsole(ret.message);
                await fetchStats();
                switchTab('tab-stats');
            } catch (e) {
                logToConsole("送出失敗: " + e);
            }
        }

        async function initSeasonRoster() {
            if (!confirm("確定要清空附件四.xlsx的參賽名冊嗎？(系統會自動備份)")) return;
            try {
                const res = await fetch(`${API_BASE}/api/init-season`, { method: 'POST' });
                const ret = await res.json();
                alert(ret.message);
                await fetchStats();
            } catch (e) {
                alert("重設失敗: " + e);
            }
        }

        function downloadStatsCSV() {
            let csv = "\uFEFF分區,編號,社名,主協辦身分,參賽人數,報名費,僅用餐人數,餐費,應收合計,匯款日期,實收金額,未收差額,狀態,備註\n";
            currentData.stats.forEach(s => {
                let status = s.received >= s.receivable && s.receivable > 0 ? "已繳款" : (s.unpaid > 0 ? "待繳款" : "-");
                csv += `"${s.zone}","${s.code}","${s.name}","${s.sponsor}",${s.players},${s.player_fee},${s.lunch},${s.lunch_fee},${s.receivable},"${s.remit_date}",${s.received},${s.unpaid},"${status}","${s.note}"\n`;
            });
            const blob = new Blob([csv], { type: "text/csv;charset=utf-8;" });
            const link = document.createElement("a");
            link.href = URL.createObjectURL(blob);
            link.download = "保齡球賽_各社報名統計.csv";
            link.click();
        }

        function downloadRosterCSV() {
            let csv = "\uFEFF球道,席位,選手姓名,性別,所屬扶輪社\n";
            currentData.lanes.forEach(s => {
                csv += `${s.lane},${s.seq},"${s.name || ''}","${s.gender || ''}","${s.club || ''}"\n`;
            });
            const blob = new Blob([csv], { type: "text/csv;charset=utf-8;" });
            const link = document.createElement("a");
            link.href = URL.createObjectURL(blob);
            link.download = "附件四_保齡球賽40道參賽名冊.csv";
            link.click();
        }

        window.addEventListener('DOMContentLoaded', () => {
            fetchStats();
            checkBackendHealth();
            setInterval(checkBackendHealth, 10000);
            setInterval(fetchStats, 30000);
        });
        fetchStats();
        checkBackendHealth();
    </script>
</body>
</html>
"""

# ==================== API 路由 ====================

@app.route("/api/health")
def api_health():
    return jsonify({
        "status": "online",
        "timestamp": datetime.now().isoformat(),
        "is_syncing": is_syncing
    })


@app.route("/")
def index():
    html_file = os.path.join(os.path.dirname(os.path.abspath(__file__)), "保齡球賽管理系統.html")
    if os.path.exists(html_file):
        with open(html_file, "r", encoding="utf-8") as f:
            return f.read()

    clubs = core.load_clubs_db()
    # 依分區分組
    zones = OrderedDict()
    for c in clubs:
        z = c["zone"]
        if z not in zones:
            zones[z] = []
        zones[z].append(c)

    today_str = datetime.now().strftime("%Y/%m/%d")
    return render_template_string(HTML_TEMPLATE, zones=zones, today_str=today_str)


@app.route("/api/stats")
def get_stats():
    """讀取 Excel 返回所有社的報名狀況與附件四排道"""
    clubs = core.load_clubs_db()

    wb = openpyxl.load_workbook(core.STATS_XLSX_PATH, data_only=True)
    sheet = wb["各社報名資料"]

    stats_list = []
    total_clubs = 0
    total_players = 0
    total_lunch = 0
    total_receivable = 0
    total_received = 0

    for r in range(2, 82):
        zone = sheet.cell(r, 1).value
        code = sheet.cell(r, 2).value
        name = sheet.cell(r, 3).value
        sponsor = sheet.cell(r, 4).value
        players = int(sheet.cell(r, 7).value or 0)
        lunch = int(sheet.cell(r, 9).value or 0)
        received = int(sheet.cell(r, 13).value or 0)

        # 贊助金計算
        sponsor_fee = 0
        if sponsor == "主辦社":
            sponsor_fee = 10000
        elif sponsor == "協辦社":
            sponsor_fee = 5000
        elif sponsor == "贊助社":
            sponsor_fee = 3000

        player_fee = sheet.cell(r, 8).value
        if player_fee is None or player_fee == 0:
            player_fee = players * 1200
        else:
            player_fee = int(player_fee)

        lunch_fee = sheet.cell(r, 10).value
        if lunch_fee is None or lunch_fee == 0:
            lunch_fee = lunch * 1000
        else:
            lunch_fee = int(lunch_fee)

        receivable = sheet.cell(r, 11).value
        if receivable is None or receivable == 0:
            receivable = sponsor_fee + player_fee + lunch_fee
        else:
            receivable = int(receivable)

        remit_date = sheet.cell(r, 12).value
        note = sheet.cell(r, 15).value
        unpaid_calc = max(0, receivable - received)

        if players > 0 or received > 0 or sponsor or lunch > 0:
            total_clubs += 1

        total_players += players
        total_lunch += lunch
        total_receivable += receivable
        total_received += received

        stats_list.append({
            "zone": zone,
            "code": code,
            "name": name,
            "sponsor": sponsor or "",
            "players": int(players),
            "player_fee": int(player_fee),
            "lunch": int(lunch),
            "lunch_fee": int(lunch_fee),
            "receivable": int(receivable),
            "remit_date": str(remit_date) if remit_date else "",
            "received": int(received),
            "unpaid": unpaid_calc,
            "note": str(note) if note else ""
        })

    roster_wb = openpyxl.load_workbook(core.ROSTER_XLSX_PATH, data_only=True)
    roster_sheet = roster_wb["參賽名單"]
    lanes_list = []

    for r in range(2, 162):
        lane = roster_sheet.cell(r, 1).value
        seq = roster_sheet.cell(r, 2).value
        p_name = roster_sheet.cell(r, 3).value
        p_gender = roster_sheet.cell(r, 4).value
        p_club = roster_sheet.cell(r, 5).value

        lanes_list.append({
            "lane": int(lane) if lane else 1,
            "seq": int(seq) if seq else 1,
            "name": p_name or None,
            "gender": p_gender or None,
            "club": p_club or None
        })

    summary = {
        "clubs_count": total_clubs,
        "players_count": int(total_players),
        "lunch_count": int(total_lunch),
        "receivable": int(total_receivable),
        "received": int(total_received),
        "unpaid": max(0, int(total_receivable - total_received))
    }

    sync_state = core.get_sync_state()

    return jsonify({
        "stats": stats_list,
        "lanes": lanes_list,
        "summary": summary,
        "sync_state": sync_state
    })


@app.route("/api/sync", methods=["POST"])
def api_sync():
    """手動觸發郵件同步 (自動依狀態推算 3 天或回溯至上次偵測時間點)"""
    global is_syncing
    if is_syncing:
        return jsonify({"message": "系統正在同步中，請稍候..."})

    is_syncing = True
    try:
        result = core.run_sync(dry_run=False)
        for l in result.get("logs", []):
            add_log(l)

        window_desc = f"{result['window_start']} 至 {result['window_end']}"
        first_tip = "【初次推算 3 天又 1 小時】" if result["is_first"] else "【回溯至上次偵測點前 1 小時】"
        msg = f"{first_tip} 同步完成！查詢區間: {window_desc}。新處理: {result['new_processed']} 封，更正異動: {result.get('updated_processed', 0)} 封，略過已處理無異動: {result.get('skipped_count', 0)} 封。"

        return jsonify({
            "message": msg,
            "logs": result.get("logs", []),
            "sync_info": result
        })
    except Exception as e:
        add_log(f"[錯誤] 手動同步失敗: {e}")
        return jsonify({"message": f"同步失敗: {e}", "logs": [str(e)]}), 500
    finally:
        is_syncing = False


@app.route("/api/simulate-registration", methods=["POST"])
def api_sim_reg():
    """模擬執秘報名輸入"""
    data = request.json or {}
    club_code = data.get("club_code")
    sponsor = data.get("sponsor") or None
    lunch_count = data.get("lunch_count") or 0
    raw_players = data.get("players") or []

    clubs = core.load_clubs_db()
    target_club = next((c for c in clubs if c["code"] == club_code), None)
    if not target_club:
        return jsonify({"message": "找不到該社"}), 404

    # 規範化參賽選手名單，確保姓名完整性 (如 PP Sure, P Lin.C夫人, Bank) 與性別純淨性
    players = []
    for p in raw_players:
        if isinstance(p, dict):
            p_name = str(p.get("name") or "").strip()
            p_nick = str(p.get("nickname") or p_name).strip()
            p_gender = str(p.get("gender") or "男").strip()

            if p_gender not in ["男", "女"]:
                if any(x in p_gender.lower() for x in ["女", "f", "female"]):
                    p_gender = "女"
                elif any(x in p_gender.lower() for x in ["男", "m", "male"]):
                    p_gender = "男"
                else:
                    # 避免前級或外部請求誤把姓名第二段當成性別 (如 name="PP", gender="Sure")
                    p_name = f"{p_name} {p_gender}".strip()
                    p_nick = p_name
                    p_gender = "女" if "夫人" in p_name else "男"

            if p_name:
                players.append({
                    "name": p_name,
                    "nickname": p_nick,
                    "gender": p_gender
                })
        elif isinstance(p, str):
            p_dict = core.parse_single_player_line(p)
            if p_dict:
                players.append({
                    "name": p_dict["name"],
                    "nickname": p_dict["nickname"],
                    "gender": p_dict["gender"]
                })

    all_registered_data = {}
    if os.path.exists(core.ROSTER_DATA_FILE):
        with open(core.ROSTER_DATA_FILE, "r", encoding="utf-8") as f:
            all_registered_data = json.load(f)

    all_registered_data[target_club["code"]] = {
        "club_name": target_club["name"],
        "tag": target_club["tag"],
        "sponsor": sponsor,
        "players": players,
        "lunch_only": [{"name": f"午餐{i+1}"} for i in range(lunch_count)],
        "updated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    }

    core.update_registration_stats_excel(
        club=target_club,
        sponsor=sponsor,
        players_count=len(players),
        lunch_count=lunch_count,
        note="模擬報名寫入"
    )

    core.update_roster_excel(all_registered_data, clubs)

    with open(core.ROSTER_DATA_FILE, "w", encoding="utf-8") as f:
        json.dump(all_registered_data, f, ensure_ascii=False, indent=2)

    return jsonify({"message": f"成功模擬 [{target_club['code']}] {target_club['name']} 報名！已同步排入附件四。"})


@app.route("/api/simulate-deposit", methods=["POST"])
def api_sim_dep():
    """模擬上海商銀入帳輸入"""
    data = request.json or {}
    club_code = data.get("club_code")
    amount = data.get("amount") or 0
    date_str = data.get("date") or datetime.now().strftime("%Y/%m/%d")
    memo = data.get("memo") or "上海商銀入帳"

    clubs = core.load_clubs_db()
    target_club = next((c for c in clubs if c["code"] == club_code), None)
    if not target_club:
        return jsonify({"message": "找不到該社"}), 404

    core.update_registration_stats_excel(
        club=target_club,
        remit_date=date_str,
        remit_amt=amount,
        note=memo
    )

    return jsonify({"message": f"成功模擬 [{target_club['code']}] {target_club['name']} 上海商銀入帳 NT$ {amount:,} 元！"})


@app.route("/api/swap-lane", methods=["POST"])
def api_swap_lane():
    """更換選手賽道席位 (支援空位移入或選手互換)"""
    data = request.json or {}
    try:
        from_lane = int(data.get("from_lane"))
        from_seq = int(data.get("from_seq"))
        to_lane = int(data.get("to_lane"))
        to_seq = int(data.get("to_seq"))
    except (TypeError, ValueError):
        return jsonify({"success": False, "message": "球道或席位號碼格式錯誤"}), 400

    try:
        res = core.swap_players_in_roster(from_lane, from_seq, to_lane, to_seq)
        add_log(f"[賽道調換] {res.get('message')}")
        return jsonify(res)
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/api/update-player", methods=["POST"])
def api_update_player():
    """手動更新指定球道席位之選手姓名、性別與所屬社"""
    data = request.json or {}
    try:
        lane = int(data.get("lane"))
        seq = int(data.get("seq"))
        name = str(data.get("name", "")).strip()
        gender = str(data.get("gender", "男")).strip()
        club = str(data.get("club", "")).strip()
    except (TypeError, ValueError):
        return jsonify({"success": False, "message": "球道或席位號碼格式錯誤"}), 400


    try:
        res = core.update_player_in_roster(lane, seq, name, gender, club)
        add_log(f"[選手修改] {res.get('message')}")
        return jsonify(res)
    except Exception as e:
        return jsonify({"success": False, "message": str(e)}), 500


@app.route("/api/init-season", methods=["POST"])
def api_init_season():
    """清空附件四"""
    core.init_season_roster()
    return jsonify({"message": "「附件四.xlsx」參賽名單已清空重設完畢！"})


@app.route("/api/reset-all", methods=["POST"])
def api_reset_all():
    """正式上線前一鍵清空所有測試資料"""
    core.reset_all_data(reset_mail_history=False)
    add_log("【正式上線重設】已清空報名統計與附件四的所有測試數據，系統恢復為全新初始狀態。")
    return jsonify({
        "success": True,
        "message": "已成功清空所有測試資料！報名統計表與附件四參賽名單已恢復為全新初始狀態。"
    })


@app.route("/download/stats")
def download_stats():
    return send_file(core.STATS_XLSX_PATH, as_attachment=True, download_name="報名統計.xlsx")


@app.route("/download/roster")
def download_roster():
    return send_file(core.ROSTER_XLSX_PATH, as_attachment=True, download_name="附件四_參賽名單.xlsx")


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5001))
    print(f"\n========================================================")
    print(f"  保齡球比賽網頁管理系統已啟動！")
    print(f"  排程模式：每 3 小時自動檢查新 Mail")
    print(f"  請使用瀏覽器開啟: http://127.0.0.1:{port}")
    print(f"========================================================\n")
    app.run(host="0.0.0.0", port=port, debug=False)
